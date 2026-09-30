from pathlib import Path
import json,zipfile,lxml.etree as E
from pptx import Presentation
root=Path(__file__).resolve().parent.parent
p=root/'.codex-build/candidate-panda-v3.pptx';prs=Presentation(p);plan=json.loads((root/'.codex-build/slide-plan-panda-v3.json').read_text());ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
errors=[];fonts=set();out=[];notes=[]
for i,s in enumerate(prs.slides,1):
 n=s.notes_slide.notes_text_frame.text
 if '[教师逐字稿]' not in n:errors.append(f'notes {i}')
 notes.append(len(n))
 if any(t in n for t in ['花朵','红花','数字花','3×4','6×8','R220','G95','B15']):errors.append(f'stale notes {i}')
 for j,sh in enumerate(s.shapes):
  if min(sh.left,sh.top)<-10 or sh.left+sh.width>prs.slide_width+10 or sh.top+sh.height>prs.slide_height+10:out.append([i,j])
  if sh.has_text_frame:
   for para in sh.text_frame.paragraphs:
    for run in para.runs:
     if run.text:fonts.add(run.font.name)
  if sh.has_table:
   for row in sh.table.rows:
    for c in row.cells:
     for para in c.text_frame.paragraphs:
      for run in para.runs:
       if run.text:fonts.add(run.font.name)
 # Both edge rails exact 8pt
 for sh,y in [(s.shapes[0],0),(s.shapes[1],prs.slide_height-101600)]:
  if abs(sh.top-y)>5 or abs(sh.height-101600)>5:errors.append(f'rail {i}')
for a,b in plan['pairs']:
 sa,sb=prs.slides[a-1],prs.slides[b-1]
 # Compare every original object's geometry. Intentional table cell answers may change text.
 for k,sh in enumerate(sa.shapes):
  other=sb.shapes[k]
  if (sh.left,sh.top,sh.width,sh.height)!=(other.left,other.top,other.width,other.height):errors.append(f'base geometry {a}-{b}/{k}')
 title1,title2=sa.shapes[2],sb.shapes[2]
 if title1.text!=title2.text:errors.append(f'title text {a}-{b}')
 if E.tostring(title1._element.find('p:txBody',ns))!=E.tostring(title2._element.find('p:txBody',ns)):errors.append(f'title style {a}-{b}')
for f in fonts:
 if not f or not f.startswith('Alibaba PuHuiTi 3.0'):errors.append(f'font {f}')
assert 500*333*24//8==499500
assert [format(v,'08b') for v in [138,179,111]]==['10001010','10110011','01101111']
assert [8*8*2,16*16*2,16*16*3]==[128,512,768]
assert [128//8,512//8,768//8]==[16,64,96]
assert [format(v,'08b') for v in [36,33,27]]==['00100100','00100001','00011011']
assert 128+8+2==138
assert [16*16*2//8,16*8*2//8,16*16*1//8,8*8*2//8]==[64,32,32,16]
assert 2**1<4<=2**2
data=json.loads((root/'.codex-build/panda-data.json').read_text())
assert data['coarse'][3]=='33201211' and set(data['coarse'][3])==set('0123')
assert '第四行的3、3、2、0、1、2、1、1' in prs.slides[13].notes_slide.notes_text_frame.text
result={'slides':len(prs.slides),'question_answer_pairs':len(plan['pairs']),'fonts':sorted(fonts),'notes_min_chars':min(notes),'out_of_bounds':out,'errors':errors,'native_table_slides':[i for i,s in enumerate(prs.slides,1) if any(sh.has_table for sh in s.shapes)]}
(root/'.codex-build/structural-panda-v3.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False));assert not errors and not out
