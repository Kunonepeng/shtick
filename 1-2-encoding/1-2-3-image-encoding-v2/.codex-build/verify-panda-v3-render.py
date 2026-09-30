from pathlib import Path
import json, hashlib, shutil, zipfile
import numpy as np
import pdfplumber
from PIL import Image, ImageDraw
from pptx import Presentation
root=Path(__file__).resolve().parent.parent
build=root/'.codex-build'; qa=build/'qa-panda-v3'; qa.mkdir(exist_ok=True)
render=Path('/private/tmp/panda-v3-qa');final=build/'candidate-panda-v3.pptx'
hashfile=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert hashfile(root/'exports/1-2-3-image-encoding-panda-v2.pptx')=='8f62d688c46823b3c8c186f019d6cef7d070e1a9525727d340bd3212a228d52d'
assert hashfile(root/'exports/1-2-3-image-encoding-panda-v1.pptx')=='110e872857ce7f577175ec84ee7e96136108ef91003148ff68ad9bbc2b956c6a'
d=json.loads((build/'panda-data.json').read_text());assert hashfile(root/d['source'])==d['sha256']
a=np.asarray(Image.open(root/d['source']).convert('RGB'));c=d['crop'];crop=a[c['y']:c['y']+c['height'],c['x']:c['x']+c['width']].astype(float);pal=np.array(d['paletteRgb'])
for n,key,colors in [(8,'coarse',4),(16,'fine',4),(16,'six',6)]:
 z=288//n;means=crop.reshape(n,z,n,z,3).mean(axis=(1,3));idx=((means[:,:,None,:]-pal[None,None,:colors,:])**2).sum(axis=-1).argmin(axis=-1)
 assert [''.join(map(str,row)) for row in idx]==d[key]
 assert np.allclose(means,d['means8' if n==8 else 'means16'])
assert a[d['pixel']['y'],d['pixel']['x']].tolist()==d['pixel']['rgb']
assert d['coarse'][3]=='33201211'
fonts=set();pdfpath=render/'candidate-panda-v3.pdf'
with pdfplumber.open(pdfpath) as pdf:
 assert len(pdf.pages)==63
 for i,p in enumerate(pdf.pages,1):
  assert p.chars, f'blank text page {i}'
  fonts.update(ch['fontname'].split('+')[-1] for ch in p.chars)
  assert all(0<=ch['x0']<=ch['x1']<=p.width+0.1 and 0<=ch['top']<=ch['bottom']<=p.height+0.1 for ch in p.chars)
assert all('AlibabaPuHuiTi' in f for f in fonts),fonts
prs=Presentation(final)
for i,s in enumerate(prs.slides,1):
 t=s.notes_slide.notes_text_frame.text
 assert '[教师逐字稿]' in t and '[问题 / 页面目的]' in t and 'undefined' not in t
for n in range(1,64):
 p=render/f'slide-{n:02}.png';assert Image.open(p).size==(1600,900);shutil.copyfile(p,qa/p.name)
shutil.copyfile(pdfpath,qa/'slides.pdf')
for k,start in enumerate(range(1,64,12),1):
 canvas=Image.new('RGB',(1600,750),'#DDDDDD');draw=ImageDraw.Draw(canvas)
 for j,n in enumerate(range(start,min(start+12,64))):
  im=Image.open(qa/f'slide-{n:02}.png').convert('RGB');im.thumbnail((400,225));x=(j%4)*400;y=(j//4)*250;canvas.paste(im,(x,y));draw.text((x+8,y+228),f'Slide {n:02}',fill='black')
 canvas.save(qa/f'montage-{k}.jpg',quality=94)
with zipfile.ZipFile(final) as z:
 media=[name for name in z.namelist() if name.startswith('ppt/media/')];unique_media={hashlib.sha256(z.read(n)).hexdigest() for n in media}
 allowed={hashfile(root/d['source'])}
 assert unique_media==allowed
changed=[]
for n in range(1,64):
 old=build/'qa-panda-v2'/f'slide-{n:02}.png'
 if not np.array_equal(np.asarray(Image.open(old).convert('RGB')),np.asarray(Image.open(qa/f'slide-{n:02}.png').convert('RGB'))):changed.append(n)
result={'candidate_sha256':hashfile(final),'bytes':final.stat().st_size,'slides':63,'source_rgb_and_all_matrices_recomputed':True,'prior_deck_and_source_unchanged':True,'notes_complete':True,'rendered_slides':63,'render_resolution':'1600x900','pdf_fonts':sorted(fonts),'rendered_text_in_bounds':True,'embedded_media_count':len(media),'unique_source_images':len(unique_media),'changed_rendered_pages_vs_v2':changed,'visual_inspection':'See recorded visual review in panda-v3-audit.md; this script verifies files and values only','native_application_check':'WPS / PowerPoint not performed'}
(build/'render-qa-panda-v3.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False))
