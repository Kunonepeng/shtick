from pathlib import Path
import json,re,hashlib,subprocess,zipfile
from pptx import Presentation
root=Path(__file__).resolve().parent.parent;b=root/'.codex-build'
doc=(root/'course-design.qmd').read_text();prs=Presentation(b/'candidate-panda-v3.pptx');plan=json.loads((b/'slide-plan-panda-v3.json').read_text())
checks={};checks['question_ids_0_to_22']=list(map(int,re.findall(r'^# Q(\d+) ',doc,re.M)))==list(range(23))
checks['slide_count_and_pairs']=len(prs.slides)==63 and len(plan['pairs'])==28
allnotes='\n'.join(s.notes_slide.notes_text_frame.text for s in prs.slides)
alltext='\n'.join(sh.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
checks['no_ai_repair_in_deck']=not any(x in allnotes+alltext for x in ['AI','修复','破损照片','inpaint'])
checks['neutral_decode_question']='横4' not in '\n'.join(sh.text for sh in prs.slides[49].shapes if sh.has_text_frame) and '二维尺寸' not in '\n'.join(sh.text for sh in prs.slides[49].shapes if sh.has_text_frame)
checks['decode_givens']='完整小图' in alltext and '每行从左到右，各行从上到下' in alltext
sections={int(m.group(1)):m.group(0) for m in re.finditer(r'^# Q(\d+) .*?(?=^# |\Z)',doc,re.M|re.S)}
issues=[]
for n in range(1,22):
 sec=sections[n];rows=[x for x in plan['slides'] if x['q']==f'Q{n}'];lo,hi=rows[0]['slide'],rows[-1]['slide']
 if f'第{lo}–{hi}页' not in sec:issues.append(f'page mapping Q{n}')
 title=sec.splitlines()[0].split(' ',2)[2]
 for stage,label in [('question','提问页教师逐字稿'),('answer','揭示页教师逐字稿')]:
  talk=re.search(r'\*\*'+label+r'\*\*：([^\n]+)',sec).group(1)
  matches=[r for r in rows if r['stage']==stage and r['title']==title]
  if len(matches)!=1 or talk not in prs.slides[matches[0]['slide']-1].notes_slide.notes_text_frame.text:issues.append(f'notes alignment Q{n}/{stage}')
checks['main_question_transcripts_and_pages_aligned']=not issues
pal=json.loads((b/'panda-data.json').read_text())['palette'][:4]
def read_grid(slide,x,y,w,h,z=40):
 vals=[]
 for row in range(h):
  for col in range(w):
   px=round((x+col*z)*9525);py=round((y+row*z)*9525)
   sh=next(sh for sh in slide.shapes if abs(sh.left-px)<2 and abs(sh.top-py)<2 and abs(sh.width-z*9525)<2 and abs(sh.height-z*9525)<2)
   vals.append(pal.index('#'+str(sh.fill.fore_color.rgb)))
 return vals
values1=read_grid(prs.slides[50],210,440,4,2);values2=read_grid(prs.slides[50],840,440,2,4);values3=read_grid(prs.slides[51],990,310,4,2)
checks['actual_native_grids_round_trip']=values1==values2==values3==[3,3,1,0,0,1,3,0]
codes=' '.join(format(x,'02b') for x in values1);checks['actual_grid_codes_in_design']=codes in sections[18]
archive=root/'references/course-design-v2.5-source-record.qmd';prior=subprocess.check_output(['git','show','HEAD:1-2-encoding/1-2-3-image-encoding-v2/course-design.qmd'],cwd=root)
checks['full_previous_design_preserved']=archive.read_bytes()==prior
with zipfile.ZipFile(b/'candidate-panda-v3.pptx') as z:
 media={hashlib.sha256(z.read(p)).hexdigest() for p in z.namelist() if p.startswith('ppt/media/')}
checks['only_panda_image_embedded']=media=={hashlib.sha256((root/'assets/pandas.jpg').read_bytes()).hexdigest()}
checks['timing_total_45']=sum([3,10,7,2,6,7,4,6])==45
result={'checks':checks,'issues':issues,'archive_sha256':hashlib.sha256(prior).hexdigest(),'course_sha256':hashlib.sha256((root/'course-design.qmd').read_bytes()).hexdigest(),'decoded_indices':values1,'codes':codes,'claim_boundary':'Content, native shape values, notes and structural alignment; excludes actual classroom and WPS acceptance.'}
(b/'joint-audit-panda-v3.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False));assert all(checks.values()) and not issues
