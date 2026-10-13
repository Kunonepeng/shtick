#!/usr/bin/env python3
"""Three successive, auditable visual refinements of genuine 70-slide GitHub v8."""
from pathlib import Path
from io import BytesIO
import hashlib,json,zipfile
import cairosvg
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[1]
SOURCE=ROOT/'1-2-3-audio-encoding-v8.pptx'
FINAL=ROOT/'1-2-3-audio-encoding-v8.1.pptx'
ppt=Presentation(SOURCE)
assert len(ppt.slides)==70
ids=[s.slide_id for s in ppt.slides]
vio=RGBColor(98,81,177); acc=RGBColor(140,100,225)
cyan=RGBColor(0,176,240);ink=RGBColor(38,38,38)
light=RGBColor(224,221,237);pale=RGBColor(246,243,255)
REG='Alibaba PuHuiTi 3.0 55 Regular';HEAVY='Alibaba PuHuiTi 3.0 115 Black'
def I(x):return Inches(x)
def q(x):return x/914400
def n(sl,content):
    sl.notes_slide.notes_text_frame.text += '\n\n【v8.1视觉引导】'+content
def tt(sh,size=None,bold=None,color=None):
    if not sh.has_text_frame:return
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if size is not None:r.font.size=Pt(size)
            if bold is not None:r.font.bold=bold
            if color:r.font.color.rgb=color
            r.font.name=HEAVY if bold else REG
def txt(s,t,x,y,w,h,size=22,color=ink,bold=False,align=None):
    sh=s.shapes.add_textbox(I(x),I(y),I(w),I(h))
    tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Pt(0)
    tf.margin_top=tf.margin_bottom=Pt(0)
    p=tf.paragraphs[0];p.text=t;p.font.name=HEAVY if bold else REG
    p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color
    if align:p.alignment=align
    return sh
def line(s,x1,y1,x2,y2,color=acc,width=2):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,I(x1),I(y1),I(x2),I(y2))
    sh.line.color.rgb=color;sh.line.width=Pt(width)
    return sh
def shape(s,x,y,w,h,color=light,fill=None):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,I(x),I(y),I(w),I(h))
    if fill:sh.fill.solid();sh.fill.fore_color.rgb=fill
    else:sh.fill.background()
    sh.line.color.rgb=color;sh.line.width=Pt(.7)
    return sh
def drop(sh):sh._element.getparent().remove(sh._element)
def audit(stage,items):
    assert len(ppt.slides)==70
    assert [s.slide_id for s in ppt.slides]==ids
    assert all(s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip() for s in ppt.slides)
    counts={}
    for s in ppt.slides:
        texts=[sh.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text.strip()]
        if texts and texts[0].startswith('练习题 '):
            counts[texts[0]]=counts.get(texts[0],0)+1
        for sh in s.shapes:
            assert -.05<q(sh.left)<13.39
            assert -.05<q(sh.top)<7.55
            assert q(sh.left+sh.width)<13.4,(stage,'right-overflow',texts[:1])
            assert q(sh.top+sh.height)<7.58,(stage,'bottom-overflow',texts[:1])
    assert len(counts)==6 and set(counts.values())=={2},counts
    if stage>=1:
        assert any(sh.shape_type==13 for sh in ppt.slides[2].shapes)
        balls=[sh for sh in ppt.slides[44].shapes if sh.shape_type==1 and .68<q(sh.width)<.75 and 2.7<q(sh.top)<3.5]
        assert len(balls)==5,len(balls)
        pos=[q(sh.left)-(.79+i*2.52) for i,sh in enumerate(balls)]
        assert all(a<b for a,b in zip(pos,pos[1:])),pos
    if stage>=2:
        assert any('二维：行×列' in sh.text for sh in ppt.slides[43].shapes if sh.has_text_frame)
    if stage>=3:
        assert all(any(sh.shape_type==1 and abs(q(sh.top)-5.48)<.02 for sh in ppt.slides[i-1].shapes) for i in [33,38,42,48,62,67])
    target=BASE/f'round-{stage}-v8.1.pptx';ppt.save(target)
    with zipfile.ZipFile(target) as z:assert z.testzip() is None
    result={'round':stage,'slides':70,'notes':70,'exercise_pairs':counts,'improvements':items,
            'sha256':hashlib.sha256(target.read_bytes()).hexdigest()}
    (BASE/f'audit-round-{stage}.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print('ROUND',stage,'PASS',target.stat().st_size)
    return target
# Round 1: semantic illustration, distinguish movement between video frames,
# make waveform/quantization curves visible from the back of the classroom.
s=ppt.slides[2]
for sh in list(s.shapes):
    if q(sh.top)>1.6:drop(sh)
pic=cairosvg.svg2png(bytestring=(BASE/'assets/sound-vibration-to-microphone.svg').read_bytes(),output_width=1800,output_height=600)
s.shapes.add_picture(BytesIO(pic),I(.84),I(1.78),width=I(11.65),height=I(3.87))
for t,x,w in [('① 音叉振动',.86,3.25),('② 空气中的声压变化传开',4.48,4.45),('③ 麦克风接收声音',9.12,3.30)]:
    txt(s,t,x,5.69,w,.57,23,acc,True)
txt(s,'注意：声波不是一条在空间里弯曲前进的线。',.88,6.50,11.65,.46,19,vio)
n(s,'视觉依据是生成的“音叉—声压疏密变化—麦克风”图；为了投影清晰采用可缩放矢量绘制。它不是声波沿正弦路径传播的图。')
s=ppt.slides[44]
balls=[sh for sh in s.shapes if sh.shape_type==1 and .49<q(sh.width)<.58 and 2.8<q(sh.top)<3.5]
assert len(balls)==5,len(balls)
for i,sh in enumerate(balls):
    sh.width=I(.72);sh.height=I(.72)
    sh.left=I(.79+i*2.52+[.13,.37,.63,.89,1.17][i]);sh.top=I(3.06)
for i in range(5):line(s,.79+i*2.52+.38,2.68,.79+i*2.52+.38,4.26,light,1.0)
n(s,'让学生比较同一小球相对于各帧相同参考线的位置，从第一帧到第五帧明显向右变化。')
s=ppt.slides[15]
for sh in s.shapes:
    if sh.shape_type==9 and q(sh.width)>7:sh.line.width=Pt(2.5);sh.line.color.rgb=acc
    if sh.has_text_frame and sh.text.strip() in [str(i) for i in range(8)]:tt(sh,23,True,vio)
n(s,'加深0至7的八条量化等级线，先数清等级再引入3 bit。')
for i in [8,11,14,15]:
    s=ppt.slides[i-1]
    for sh in s.shapes:
        if sh.shape_type==9 and .07<q(sh.width)<.15 and q(sh.height)<.25:
            sh.line.width=Pt(2)
    n(s,'加强波形曲线投影对比；灰色线为信号示意，离散点是取样时刻的记录。')
audit(1,['P0：第45页小球的相对位移太隐蔽','P1：第3页三文字框未把物理声波直观化','P1：第16页等级线过淡','P2：若干采样曲线细且低对比'])
# Round 2: dense data and comparison typographic hierarchy.
s=ppt.slides[21]
for sh in s.shapes:
    if sh.has_text_frame and sh.text.strip() in ['采样顺序','量化等级','3 bit码字']:tt(sh,23,True,vio)
    elif sh.has_text_frame and sh.text.strip() in ['000','001','010','011','100','101','110','111']:tt(sh,20,True,cyan)
    elif sh.shape_type==9:sh.line.width=Pt(1.6)
n(s,'码字行放大并统一青色；先指出第1列，再读完全部10个样本。')
s=ppt.slides[43]
for sh in list(s.shapes):
    if 1.7<q(sh.top)<5.9:drop(sh)
txt(s,'比较',.83,1.75,2.33,.6,24,vio,True)
txt(s,'图像：空间采样',3.42,1.75,4.18,.6,26,vio,True)
txt(s,'声音：时间采样',8.07,1.75,4.34,.6,26,cyan,True)
for i,(a,b,c) in enumerate([
 ('① 取样位置','二维：行×列的像素','一维：时间点'),
 ('② 记录什么','像素颜色','信号幅度'),
 ('③ 数据量怎么算','像素数 × 每像素位数','样本数 × 每样本位数')]):
    y=2.71+i*1.20;line(s,.84,y-.17,12.31,y-.17,light,1.2)
    txt(s,a,.85,y,2.48,.7,22,vio,True)
    txt(s,b,3.43,y,4.12,.77,22)
    txt(s,c,8.08,y,4.35,.77,22)
n(s,'重新排列三条对照，文字达到22pt，学生先指出“位置”不同，再比较数据量结构。')
s=ppt.slides[58]
for sh in s.shapes:
    if not sh.has_text_frame:continue
    val=sh.text.strip()
    if val in ['ZIP / RAR','MP3','WAV','MP4']:tt(sh,30,True,vio if val in ['ZIP / RAR','WAV'] else cyan)
    elif val in ['通常为无损归档格式','常见有损音频编码','容器；常见未压缩PCM','容器；常装有损视频']:tt(sh,22,False,ink)
n(s,'大字号提示学生区分归档格式、压缩编码与文件容器；颜色是视觉分组而非“都同一类”。')
s=ppt.slides[49]
for sh in s.shapes:
    if not sh.has_text_frame:continue
    if sh.text.strip() in ['未压缩图像帧','重新编码后的视频']:tt(sh,25,True,vio)
    if sh.text.strip() in ['每一帧都要保存完整的原始像素数据','利用帧间相似和更有效的编码，数据可能显著减小']:tt(sh,22,False,ink)
n(s,'将关键概念和对照放大；不能将视频压缩简化为“直接删除几帧”。')
audit(2,['P1：10个3-bit码字偏小','P1：图像/声音比较表格层级低','P1：格式分类过小','P2：未压缩与压缩对照偏弱'])
# Round 3: consistent question/answer reveal emphasis, not altering the question positions.
for idx in [33,38,42,48,62,67]:
    s=ppt.slides[idx-1]
    bg=shape(s,.85,5.48,11.66,1.38,light,pale)
    s.shapes._spTree.remove(bg._element);s.shapes._spTree.insert(4,bg._element)
    bar=shape(s,.85,5.48,.075,1.38,vio,vio)
    s.shapes._spTree.remove(bar._element);s.shapes._spTree.insert(5,bar._element)
    for sh in s.shapes:
        if not sh.has_text_frame:continue
        if sh.text.strip().startswith('答案：'):
            sh.left=I(1.10);sh.width=I(11.20)
            sh.top=I(5.59);sh.height=I(.57);tt(sh,26,True,vio)
        elif 6.1<q(sh.top)<6.7 and sh.text.strip():
            sh.left=I(1.10);sh.width=I(11.20);tt(sh,18,False,ink)
    n(s,'题干完全保留原位置，仅答案区浅紫底原位揭示；注意先等待学生回应再翻页。')
s=ppt.slides[64]
for sh in s.shapes:
    if not sh.has_text_frame:continue
    if sh.text.strip() in ['重要文档长期备份','实时视频通话']:
        tt(sh,27,True,vio if sh.text.strip().startswith('重要') else cyan)
    elif sh.text.strip().startswith('优先'):tt(sh,23)
n(s,'两组场景对比加大；让学生按完整性、带宽、延迟、质量权衡。')
result=audit(3,['P1：6组答案视觉强调不足','P1：两种压缩场景字号较小','P2：70页逐页边界、讲者备注和题答成对复核'])
FINAL.write_bytes(result.read_bytes())
print('PUBLISHED',FINAL,FINAL.stat().st_size)
