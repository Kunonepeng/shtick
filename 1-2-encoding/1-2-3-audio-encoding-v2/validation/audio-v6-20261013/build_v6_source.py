#!/usr/bin/env python3
"""Audio encoding & data compression v6: two 45-minute Chinese classroom lessons.
Transcripts-based pedagogical reconstruction of v5 topics with native diagrams.
Source references: three teacher transcripts; PPT visual guide; curriculum standard.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pathlib import Path
from math import sin, pi
import math, random, zlib

ROOT=Path(__file__).resolve().parents[2] if 'validation' in __file__ else Path('/mnt/data/v6_work')
OUTPUT=ROOT/'1-2-3-audio-encoding-v6.pptx'
p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
V=RGBColor(98,81,177); A=RGBColor(140,100,225); C=RGBColor(0,176,240); R=RGBColor(255,0,0)
BL=RGBColor(38,38,38); GR=RGBColor(128,128,128); LG=RGBColor(226,226,231); WHITE=RGBColor(255,255,255); BLACK=RGBColor(0,0,0)
REG='Alibaba PuHuiTi 3.0 55 Regular'; BOLD='Alibaba PuHuiTi 3.0 115 Black'; GREEN=RGBColor(26,141,93)
GRAY_BG=RGBColor(248,248,251)

SLIDEDATA=[]
def box(s,x,y,w,h,color=None,border=None,bw=1):
  sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
  if color: sh.fill.solid();sh.fill.fore_color.rgb=color
  else: sh.fill.background()
  if border: sh.line.color.rgb=border;sh.line.width=Pt(bw)
  else: sh.line.fill.background()
  return sh

def tb(s,txt,x,y,w,h,pt=22,bold=False,color=BL,align=None):
  sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));f=sh.text_frame;f.clear();f.word_wrap=True
  f.margin_left=f.margin_right=Pt(0);f.margin_top=f.margin_bottom=Pt(0)
  for i,t in enumerate(str(txt).split('\n')):
    q=f.paragraphs[0] if i==0 else f.add_paragraph();q.text=t;q.space_before=Pt(0);q.space_after=Pt(2)
    q.font.name=BOLD if bold else REG;q.font.size=Pt(pt);q.font.bold=bold;q.font.color.rgb=color
    if align is not None:q.alignment=align
  return sh

def line(s,x,y,xx,yy,col=LG,w=2):
  sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x),Inches(y),Inches(xx),Inches(yy));sh.line.color.rgb=col;sh.line.width=Pt(w);return sh

def circle(s,x,y,d,color=WHITE,border=A,bw=2):
  sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=color;sh.line.color.rgb=border;sh.line.width=Pt(bw);return sh

def slide(title,notes,lesson=1,kind='content'):
  s=p.slides.add_slide(p.slide_layouts[6]);box(s,0,.04,13.333,.08,V);box(s,0,7.39,13.333,.08,V)
  if kind!='cover':tb(s,title,.67,.63,12,.86,34,True)
  s.notes_slide.notes_text_frame.text=notes
  SLIDEDATA.append((len(p.slides),lesson,title))
  return s

def caption(s,txt,y=6.43,color=A):tb(s,txt,.85,y,11.8,.50,20,True,color)

def explain(s,a,b=None,x=.89,y=2.05,w=11.7,pt=27):
  tb(s,a,x,y,w,1.0,pt,True,A)
  if b:tb(s,b,x,y+1.24,w,1.80,22)

def diagram_label(s,label,x,y,w,col=A,pt=22,stroke=True):
  if stroke:box(s,x,y,w,.73,None,col,2)
  tb(s,label,x+.08,y+.13,w-.16,.44,pt,True,col,PP_ALIGN.CENTER)

def step_arrow(s,x,y,xx,yy,col=A):line(s,x,y,xx,yy,col,3)

def question(n,txt,student_prompt,teacher_notes,lesson=1,source=''):
  s=slide('练习题 '+str(n), '问题页。'+teacher_notes,lesson)
  tb(s,txt,.86,1.85,11.65,3.32,27,False)
  line(s,.90,5.38,12.12,5.38,LG,2)
  tb(s,student_prompt,.91,5.65,11.3,.74,22,True,A)
  return s

def reveal(n,txt,answer,reason,teacher_notes,lesson=1):
  s=slide('练习题 '+str(n), '答案页。'+teacher_notes,lesson)
  tb(s,txt,.86,1.85,11.65,3.32,27,False)
  line(s,.90,5.38,12.12,5.38,LG,2)
  tb(s,answer,.91,5.55,11.5,.55,23,True,A)
  tb(s,reason,.91,6.19,11.5,.66,19,False)
  return s

def axes(s,x,y,w,h,title='',ylabel='振幅',xlabel='时间'):
  line(s,x,y+h*.5,x+w,y+h*.5,GR,1.6);line(s,x+.10,y+.04,x+.10,y+h-.03,GR,1.3)
  tb(s,ylabel,x+.13,y-.30,1.4,.35,16,False,GR)
  tb(s,xlabel,x+w-1,y+h*.5+.13,.9,.32,16,False,GR)
  if title:tb(s,title,x,y-.55,w,.45,22,True,A)

def wave(s,x,y,w,h,frequency=2,amp=.37,color=A,stroke=2.6,n=100,phase=0):
  coords=[]
  for i in range(n+1):
    t=i/n;px=x+t*w;py=y+h/2 - amp*h*math.sin(2*math.pi*frequency*t+phase)
    coords.append((px,py))
  for a,b in zip(coords[:-1],coords[1:]):line(s,a[0],a[1],b[0],b[1],color,stroke)
  return lambda t:y+h/2-amp*h*math.sin(2*math.pi*frequency*t+phase)

def sample_points(s,x,y,w,h,n=9,frequency=2,amp=.37,color=C,connect=False,guide=False,phase=0):
  coords=[]
  for i in range(n):
    t=i/n;xx=x+t*w;yy=y+h/2-amp*h*math.sin(2*math.pi*frequency*t+phase)
    coords.append((xx,yy))
    if guide:line(s,xx,y+.03,xx,y+h-.01,LG,1)
    circle(s,xx-.067,yy-.067,.134,color, color,1)
  if connect:
    for a,b in zip(coords[:-1],coords[1:]):line(s,*a,*b,C,2)
  return coords

def banner(s,lead,rest='',y=6.03):
  tb(s,lead,.89,y,10.9,.52,23,True,A)
  if rest:tb(s,rest,.92,y+.59,11.6,.42,19)

def compare(s,left_head,left_body,right_head,right_body,footer=None):
  line(s,6.65,1.78,6.65,5.91,LG,2)
  tb(s,left_head,.90,1.90,5.55,.76,28,True,A);tb(s,right_head,7.06,1.90,5.2,.76,28,True,C)
  tb(s,left_body,.90,3.00,5.48,2.50,23);tb(s,right_body,7.06,3.00,5.33,2.50,23)
  if footer:caption(s,footer)

# ───── 第一课时：声音数字化 ─────
s=slide('第1课时：声音怎样变成0和1？', '开场。问：“上一节课学习了图像编码。那么在手机里录下我们的声音，计算机保存的到底是什么？”不要直接报出采样、量化、编码，先让学生提出问题。',1,'cover')
tb(s,'声音怎样变成0和1？',.75,2.25,11.9,1.27,52,True)
tb(s,'第一课时｜从声音到数字音频',.80,4.05,10.9,.82,31,False,GR)

s=slide('我们会编码文字、图像，那声音呢？','回忆旧知识：字符编码记录符号，图像编码记录像素。请学生预测：自然声音没有像素，计算机怎么办？指向本课总问题，不求答对。')
diagram_label(s,'文字 → 编码',.86,2.25,3.5);diagram_label(s,'图像 → 编码',4.95,2.25,3.5,C)
tb(s,'声音 → ？',9.24,2.20,3.3,.9,32,True,A)
caption(s,'核心问题：手机为什么可以把正在说的话保存成文件？',5.45)

s=slide('声音首先是什么？','从实录1的振动声源、空气传播与耳膜经验出发。问：我们说话时什么振动？声音怎样到耳朵？要学生识别声音是物理现象。')
for i,(name,txt) in enumerate([('声带振动','发出声音'),('空气传播','声压变化传开'),('耳膜振动','听到声音')]):
  x=.85+i*4.20;diagram_label(s,name,x,2.50,3.25,A if i!=1 else C)
  tb(s,txt,x,3.60,3.25,.58,22,False,BL,PP_ALIGN.CENTER)
  if i<2:step_arrow(s,x+3.34,2.86,x+4.08,2.86)
caption(s,'声音是一种在介质中传播的机械波。',5.55)

s=slide('声音和电脑里的音频文件，是同一种量吗？','先让学生表态：声波在时间上连续变化；文件保存的是离散数字。问麦克风前后，哪一阶段发生改变？')
compare(s,'自然界的声音','空气中的声压随时间连续变化。\n\n这是模拟量。','电脑中的音频文件','计算机记录一串数值。\n\n这是数字量。','我们的任务：让连续的声音变成计算机能够记录的数字。')

s=slide('麦克风把声波变成什么？','承接实录1：猜需要哪些设备。麦克风将声压变化转为连续电信号；此时仍是模拟信号。不要说麦克风直接保存WAV。')
diagram_label(s,'空气中的声音',.82,2.5,3.4);step_arrow(s,4.36,2.87,5.20,2.87)
diagram_label(s,'麦克风',5.25,2.5,2.55,C);step_arrow(s,7.93,2.87,8.82,2.87)
diagram_label(s,'连续电信号',8.88,2.5,3.25)
caption(s,'麦克风完成了“声波 → 电信号”，但信号仍然连续。',5.38)

s=slide('连续电信号，怎样变成数字？','先问：谁来把模拟电信号转换成0和1？鼓励学生猜CPU、麦克风等。展示声卡音频接口或画流程，以ADC为答案。设备可集成在主板或外置。')
diagram_label(s,'模拟电信号',.88,2.5,3.50);step_arrow(s,4.52,2.87,5.18,2.87)
diagram_label(s,'声卡：ADC',5.28,2.5,3.0,C);step_arrow(s,8.43,2.87,9.1,2.87)
diagram_label(s,'数字音频',9.2,2.5,3.0)
tb(s,'ADC = 模数转换（模拟信号 → 数字信号）',1.05,4.45,11.5,.84,25,True,A)
caption(s,'声卡负责完成这一转换；它不等于“存储音频文件”。',5.67)

s=slide('把录音过程和播放过程接起来','追问：存进去的0和1，为什么还能从扬声器变回声音？给出DAC作为反向转换，重点是两个方向的理解；不要展开电路。')
for i,name in enumerate(('声音','麦克风','ADC','数字文件')):
  x=.7+i*3.12;diagram_label(s,name,x,2.23,2.52,C if i==2 else A)
  if i<3:step_arrow(s,x+2.58,2.59,x+3.07,2.59)
for i,name in enumerate(('数字文件','DAC','扬声器','声音')):
  x=.7+i*3.12;diagram_label(s,name,x,4.25,2.52,C if i==1 else A)
  if i<3:step_arrow(s,x+2.58,4.61,x+3.07,4.61)
tb(s,'录音（ADC）',.84,1.66,4.3,.45,21,True,A)
tb(s,'播放（DAC）',.84,3.68,4.3,.45,21,True,A)

s=slide('波形图画的是“声音的传播路线”吗？','学生易误解：上下起伏是空气中声波的传播轨迹。请先判断，再明确横轴时间、纵轴同一地点声压或电信号的变化。')
axes(s,1.1,2.23,10.8,3.15,'同一个位置，信号随时间变化')
wave(s,1.25,2.24,10.0,3.05,frequency=2.1)
caption(s,'波形不是声音在空间里走过的弯曲路线。')

s=slide('如果只能记录有限个数字，应该记录哪些时刻？','先在波形上让学生指出选取若干时刻的方法，预测“隔一段时间取一个值”；不要急着宣布定义。')
axes(s,1.0,2.1,11.0,3.42)
wave(s,1.17,2.11,10.23,3.23,frequency=1.85)
caption(s,'请提出一个规则：怎样公平地选出一组采样时刻？')

s=slide('每隔相同时间，取一个幅度值。','保留前页坐标与波形，叠加等间隔采样点。追问：采样之后还保留每个时刻的原始值吗？概括“时间上离散”。')
axes(s,1.0,2.1,11.0,3.42)
wave(s,1.17,2.11,10.23,3.23,frequency=1.85,color=GR,stroke=1.8)
sample_points(s,1.17,2.11,10.23,3.23,n=12,frequency=1.85,guide=True)
caption(s,'采样：在固定时刻取样，让时间上的记录变得离散。')

s=slide('1 Hz和2 Hz，一秒内取几个样本？','先让学生预测1 Hz、2 Hz分别表示每秒几个点；同一时间范围，不同点数。讲采样率=单位时间样本数，间隔Δt=1/fs。')
for j,(lab,n) in enumerate((('1 Hz',4),('2 Hz',8))):
  y=1.95+j*2.15
  axes(s,1.04,y,10.8,1.52,lab)
  sample_points(s,1.25,y,9.9,1.46,n=n,frequency=1.0,amp=.35,color=A if j==0 else C)
caption(s,'采样率：每秒采集多少个样本，单位为Hz。',6.35)

s=slide('提高采样率，改变的是时间轴上的什么？','同一波形、同一个时长，采样点由少到多。先比较时间方向的点密度，不谈“采样后连线就完全恢复”。')
for j,(lab,n) in enumerate((('采样点较少',8),('采样点较多',20))):
  y=1.86+j*2.30
  axes(s,1.10,y,10.66,1.60,lab)
  wave(s,1.23,y,9.96,1.57,frequency=1.8,color=GR,stroke=1.6)
  sample_points(s,1.23,y,9.96,1.57,n=n,frequency=1.8,color=A if j==0 else C)
caption(s,'每秒记录的次数变多；数据量也会随之增加。',6.52)

s=slide('采样点的值，可能是3.71这样的数。','实录2第二步：采样点的振幅可在连续范围取值，设备却计划使用有限等级。让学生意识到有第二个问题尚未解决。')
axes(s,.95,2.10,11.10,3.12)
wave(s,1.1,2.12,10.22,3.05,frequency=1.5,color=GR,stroke=1.6)
coords=sample_points(s,1.1,2.12,10.22,3.05,n=10,frequency=1.5)
if coords:
  x,y=coords[3];circle(s,x-.13,y-.13,.26,WHITE,R,3);tb(s,'某个样本：3.71',x+.18,y-.44,3.15,.53,20,True,A)
caption(s,'新问题：如果只允许8个振幅等级，怎么记录这些值？')

s=slide('把振幅划分成0～7，共几个等级？','坚持实录2的首个量化模型：等级从0数到7含0，共八级。提醒这里是教学简化例，真实PCM可正可负。让学生先自己数等级。')
for i in range(8):
  yy=5.90-i*.53
  line(s,1.75,yy,10.50,yy,LG if i%2 else A,1.5)
  tb(s,str(i),.89,yy-.18,.65,.40,19,True,A)
caption(s,'0也是一个等级，所以一共有8个等级。')

s=slide('量化：把每个样本归入一个代表值。','此页展示明确的取整规则：教学用向下取整（非真实所有PCM的唯一规则）；把3.71映射到3。解释损失了细节，产生量化误差。')
for i in range(8):
  yy=5.75-i*.50;line(s,2.08,yy,9.8,yy,LG,1.5);tb(s,str(i),1.18,yy-.18,.7,.4,19,True,A)
# sample at 3.71: level y(3.71), representative y(3)
y0=5.75-3.71*.50;yq=5.75-3*.50
circle(s,6.1,y0-.095,.19,WHITE,R,3);circle(s,6.1,yq-.095,.19,C,C,2)
line(s,6.20,y0+.11,6.20,yq-.11,R,2.7)
tb(s,'原值 3.71',6.65,y0-.45,3.0,.46,21,True)
tb(s,'记录 3',6.65,yq+.04,2.5,.47,21,True,A)
caption(s,'量化：用有限个等级来近似表示样本的振幅。')

s=slide('一个样本需要几个二进制位？','出选择：8个等级需要2位、3位还是8位？请学生复习2的n次方。按最简单映射列出对应表，避免一次上来10个样本。')
tb(s,'8个等级 = 2³',.91,1.84,7.2,.77,31,True,A)
for i in range(8):
  x=1.00+(i%4)*3.05;y=2.98+(i//4)*1.30
  diagram_label(s,f'{i} → {i:03b}',x,y,2.48,A if i<4 else C,23)
caption(s,'3 bit可以写出8种不同的编码：000～111。')

s=slide('同一个声音的10个样本，怎样依次编码？','引导从单样本迁移到样本序列：强调时间先后、量化值与3位二进制对应。不要把10个样本误说为10个量化等级。')
vals=[3,5,7,6,4,2,0,1,3,4]
tb(s,'采样顺序',.78,2.02,2.0,.48,19,True,A)
tb(s,'量化等级',.78,3.08,2.0,.48,19,True,A)
tb(s,'3 bit码字',.78,4.13,2.0,.48,19,True,A)
for i,v in enumerate(vals):
  x=2.94+i*.94
  tb(s,str(i+1),x,2.03,.80,.42,18,False,GR,PP_ALIGN.CENTER)
  tb(s,str(v),x,3.07,.80,.43,23,True,A,PP_ALIGN.CENTER)
  tb(s,f'{v:03b}',x,4.13,.80,.43,19,True,C,PP_ALIGN.CENTER)
  line(s,x+.39,2.46,x+.39,2.97,LG,1)
  line(s,x+.39,3.49,x+.39,4.03,LG,1)
caption(s,'采样决定“什么时候记录”，量化决定“记录哪个等级”，编码决定“用哪些bit表示”。')

s=slide('把量化等级分得更细，会发生什么？','引入实录2的尺子类比：原本分米刻度与厘米刻度，刻度更细通常误差更小。与采样率的时间密度作区分。')
for j,(lab,n) in enumerate((('3 bit → 8级',8),('4 bit → 16级',16))):
  x=1.5+j*6.0;tb(s,lab,x,1.92,4.75,.56,26,True,A if j==0 else C)
  for i in range(n):
    yy=5.72-i*(3.08/(n-1))
    line(s,x,yy,x+4.45,yy,LG if i%2 else (A if j==0 else C),1)
  tb(s,'可选代表值更密集',x,6.0,4.5,.49,20)

s=slide('不要混淆：采样率和量化位数。','学生两两互问：采样率改变什么方向？量化位数改变什么方向？画时间与幅度两个方向的粗箭头，对照两种概念。')
axes(s,1.06,2.1,10.7,3.0)
wave(s,1.21,2.12,10.0,2.90,frequency=1.8,color=GR,stroke=1.8)
sample_points(s,1.21,2.12,10.0,2.90,n=13,frequency=1.8,color=C)
line(s,4.4,5.20,7.9,5.20,C,4);tb(s,'采样率：时间方向',4.0,5.51,4.6,.55,23,True,C)
line(s,11.62,2.30,11.62,4.92,A,4);tb(s,'量化位数\n幅度方向',9.95,2.05,2.3,1.14,22,True,A)

s=slide('采样更多、等级更密，就一定听得更好吗？','技术纠偏：在其他条件一致且原先不足时通常改善还原保真度；但达到合理条件后并非“数字越大越好”，设备、声源、听觉与滤波都会限制。中学只需掌握趋势和条件。')
tb(s,'采样率 ↑',1.00,2.40,4.60,.77,30,True,A)
tb(s,'每秒样本更多',1.00,3.27,4.8,.70,23)
tb(s,'量化位数 ↑',7.10,2.40,4.70,.77,30,True,C)
tb(s,'每样本可选等级更多',7.10,3.27,5.0,.70,23)
banner(s,'通常能提高记录的精细程度，但音质不会无限提高。',y=5.72)

s=slide('3 bit是教学模型，真实录音常见16 bit。','避免从课堂3位例子跳到真实16位时困惑。说明用3位为了看清规律，不是说实际音乐普遍使用3位。')
diagram_label(s,'课堂模型：3 bit',.99,2.26,5.15,A,26)
diagram_label(s,'PCM例子：16 bit',7.14,2.26,5.06,C,26)
for j,n in enumerate((3,16)):
  x=1.1 if j==0 else 7.23
  for i in range(n):
    xx=x+(i%8)*.53; yy=3.75+(i//8)*.64
    box(s,xx,yy,.42,.45,WHITE,A if j==0 else C,2)
    tb(s,'0' if i%3 else '1',xx+.10,yy+.035,.25,.32,16,True)
caption(s,'位数越多，一份未压缩样本数据所占的空间也越大。')

s=slide('PCM样本和WAV文件，是一回事吗？','从实录与v4延展：PCM代表样本数据的编码表示，WAV是保存音频数据的文件容器之一。暂不讲复杂格式比较，示意文件头告诉播放器如何读数据。')
box(s,1.0,2.38,3.17,1.5,None,A,2);tb(s,'WAV文件头',1.26,2.81,2.7,.5,25,True,A)
box(s,4.36,2.38,7.72,1.5,None,C,2);tb(s,'PCM：一个个数字样本',4.85,2.81,6.86,.50,25,True,C)
tb(s,'采样率、声道等说明',1.02,4.31,4.0,1.04,20)
tb(s,'例如：0010 1101 0110 …',4.52,4.32,7.5,1.04,22)
caption(s,'编码 ≠ 压缩；WAV常保存未压缩PCM，但并非只有这一种情况。')

s=slide('播放时：数字样本怎样重新变成声音？','回扣开场录音/播放闭环：读PCM样本→转换为连续电信号→推动扬声器振动；波形经过重建，不是简单用折线连接采样点。')
for i,(lab,col) in enumerate((('数字样本',A),('DAC',C),('连续电信号',A),('扬声器',C))):
  x=.77+i*3.2;diagram_label(s,lab,x,2.47,2.63,col)
  if i<3:step_arrow(s,x+2.71,2.84,x+3.10,2.84)
caption(s,'录音是模拟 → 数字；播放是数字 → 模拟。',5.3)

s=slide('声音数据量为什么与四个参数有关？','从实录2由学生提出因素：越长录音越大，每秒越多点越大，每个点越多bit越大，左右两路各记录一次。请先预测四项，不马上展示公式。')
for i,(head,desc) in enumerate((('采样率','每秒几个样本'),('量化位数','每样本多少bit'),('声道数','同一时刻几路'),('时长','持续多少秒'))):
  x=.9+(i%2)*6.06;y=2.06+(i//2)*1.78
  tb(s,head,x,y,5.4,.48,26,True,A if i%2==0 else C)
  tb(s,desc,x,y+.64,5.4,.74,21)

s=slide('先数样本，再数bit，最后换算成Byte。','用小模型实推：每秒2个样本×4秒×1声道=8个样本，每样本3bit共24bit=3Byte。用8个小方块表示样本，而不是一上来给公式。')
tb(s,'2 样本/秒 × 4 秒 × 1 声道 = 8 个样本',.9,1.90,11.8,.70,25,True)
for i in range(8):
  x=1.03+i*1.41;box(s,x,3.02,1.12,.84,None,A,2);tb(s,'3 bit',x+.13,3.28,.98,.36,18,True,A)
tb(s,'8 个样本 × 3 bit = 24 bit = 3 Byte',.91,4.80,11.7,.64,27,True,C)
caption(s,'不同声道要各自记录样本；声道数翻倍，未压缩样本数据量也翻倍。')

s=slide('未压缩PCM数据量公式，从哪里来？','不要把数据量公式当背诵口令。请学生逐个解释每一项代表哪个物理意义。单位：Hz、bit、秒；/8得到字节。')
tb(s,'字节数 = 采样率 × 位数 × 声道数 × 秒数 ÷ 8',.84,2.15,11.7,1.12,29,True,A)
line(s,.92,3.56,12.23,3.56,LG,2)
tb(s,'Hz：每秒采多少次',1.0,3.94,5.1,.60,22)
tb(s,'bit：每个样本用几位',6.78,3.94,5.2,.60,22)
tb(s,'声道：有几组数据',1.0,4.94,5.1,.60,22)
tb(s,'秒：记录多久',6.78,4.94,5.2,.60,22)
caption(s,'1 kHz=1000 Hz；8 bit=1 Byte。',6.20)

q1='采样率 16,000 Hz，8 bit，双声道，录制5秒。\n未压缩PCM样本数据共有多少 Byte？'
question(1,q1,'先标出每个数字的含义。', '给学生独立列式。若学生写16×8×2×5，追问16kHz应等于多少Hz。',1)
reveal(1,q1,'答案：160,000 Byte','16,000 × 8 × 2 × 5 ÷ 8 = 160,000 Byte。','问：如果双声道改成单声道，结果会怎样？',1)

s=slide('本节课你能解释三件事吗？','第一课时退出评价。随机请学生说出完整录音链；指一个点说采样与量化区别；算一条小PCM数据量。没掌握则下节开头再复习，不拖到所有附录。')
for i,(q,a) in enumerate((('① 声音怎么进入电脑？','麦克风 → ADC → 数字数据'),('② 采样和量化有什么不同？','前者定时刻，后者定代表值'),('③ PCM为什么有大小差别？','样本数 × 每样本bit'))):
  y=1.95+i*1.42;tb(s,q,.88,y,11.8,.49,24,True,A);tb(s,a,1.05,y+.62,11.5,.52,21)

# ───── 第二课时：数据量、视频、压缩 ─────
s=slide('第2课时：为什么音频、视频还需要压缩？','第二课时导入。说明第一节学的是数字化与PCM数据量，第二节学数据量计算、音轨和声道、视频数据量冲突及压缩。',2,'cover')
tb(s,'为什么文件还需要压缩？',.76,2.18,11.9,1.32,50,True)
tb(s,'第二课时｜音频、视频编码与数据压缩',.79,4.08,11.7,.85,29,False,GR)

s=slide('先回忆：这四个数字，分别是什么意思？','实录3起手计算前，先让学生能解释参数。分组对答48,000Hz、16bit、2声道、60秒。',2)
for i,(a,b) in enumerate((('48,000 Hz','每秒取样次数'),('16 bit','每样本的位数'),('2','左右两个声道'),('60 s','音频持续时间'))):
  x=.94+(i%2)*6.12;y=2.00+(i//2)*1.90
  tb(s,a,x,y,5.42,.62,30,True,A if i%2==0 else C)
  tb(s,b,x,y+.79,5.42,.57,22)
caption(s,'提醒：1 MB（按1024²换算）在严格单位写法中应称MiB。')

q2='1分钟未压缩PCM音频：48,000 Hz、16 bit、双声道。\n样本数据量约多少MiB？'
question(2,q2,'先列式，再做单位换算。','复现实录3的DVD参数题，强调对参数意义的解释，而不只是按键计算。',2)
reveal(2,q2,'答案：约10.99 MiB','48,000×16×2×60÷8=11,520,000 B；÷1024²≈10.99 MiB。','实录转写常写10.98MB，这是截断近似；按四舍五入保留两位应10.99 MiB。',2)

q3='给5分钟的诗朗诵加入一条等时长背景音乐音轨。\n如果导出时仍为44.1 kHz、16 bit、双声道PCM，\n数据量会翻倍吗？'
question(3,q3,'先投票：翻倍 / 不变 / 说不清。', '先让学生预测。这是实录3较长的认知冲突案例，不要马上公布答案。',2)

s=slide('两个看似有道理的答案，你支持哪一个？','把实录中两位学生的理由换一种说法重现。引导学生指出分歧其实在“音轨和声道是不是一回事”。仍不公布答案。',2)
compare(s,'观点 A：不变','采样率没变，位数没变，时长没变，输出的声道数也没变。','观点 B：翻倍','录音里多了一份音乐，是不是就需要多一份音频数据？','争论的关键：编辑时的“音轨”，是不是导出时的“声道”？')

s=slide('音轨是编辑时的层，声道是输出时的通道。','直观图示：两条音轨经混音成为两个输出声道，注意不是左右各有一条编辑音轨。解释可单声道或双声道输出。',2)
for i,(label,col) in enumerate((('音轨1：诗朗诵',A),('音轨2：背景音乐',C))):
  y=2.35+i*1.02;box(s,.94,y,4.12,.75,None,col,2);tb(s,label,1.11,y+.15,3.82,.44,21,True,col)
step_arrow(s,5.37,3.22,6.93,3.22)
tb(s,'混音导出',5.59,2.55,2.1,.42,20,True,A)
for i,label in enumerate(('左声道 L','右声道 R')):
  y=2.35+i*1.02;box(s,7.40,y,4.62,.75,None,C,2);tb(s,label,7.73,y+.15,4.00,.44,21,True,C)
caption(s,'编辑有两条音轨，导出仍然可以只有两个声道。',5.68)

reveal(3,q3,'答案：未压缩样本数据量不变。','内容变了，但采样率、位数、声道数、时长均未变。','让学生回看第一个预测，承认纠正了哪一个概念。提醒不同格式和压缩参数会影响最终文件大小。',2)

s=slide('迁移：把双声道改成单声道，会怎样？','让学生应用四因素公式。声道数2→1，其他不变，PCM样本数据量减半。音轨增减并不直接决定总量。',2)
tb(s,'双声道：2路数据',.96,2.39,5.4,.60,27,True,A)
tb(s,'单声道：1路数据',7.11,2.39,5.4,.60,27,True,C)
for i in range(2):box(s,1.22,3.44+i*.63,4.04,.45,None,A,2)
box(s,7.36,3.76,4.04,.45,None,C,2)
caption(s,'未压缩PCM：声道数变成一半，样本数据量也变成一半。')

s=slide('图像采样与声音采样，哪里相同？','实录3转向视频前回顾图像和声音：图像在空间二维网格采样，声音在时间轴取值；量化都用有限等级记录。',2)
for i,(a,b,c) in enumerate((('取样位置','二维：行×列的像素','一维：时间点'),('描述的量','像素的颜色信息','声音信号振幅'),('数据量','像素数×每像素位数','样本数×每样本位数'))):
  y=2.07+i*1.22;tb(s,a,.82,y,2.53,.60,22,True,A);tb(s,b,3.46,y,4.40,.70,21);tb(s,c,8.10,y,4.15,.70,21)
  if i<2:line(s,.9,y+.88,12.2,y+.88,LG,1.2)
caption(s,'这两种编码都有“选择样本”和“用有限数字表达样本”的过程。')

s=slide('一张张图片连续变化，就能让我们看到运动。','借鉴实录3翻页小人的经验。每一帧呈现同一圆的位置改变。问学生：视频需保存几个画面？帧率如何描述？不用严格视觉暂留时间阈值。',2)
for i in range(5):
  x=.79+i*2.52;box(s,x,2.28,2.26,2.28,None,LG,1.5)
  circle(s,x+.20+.28*i,3.16,.53,C,C,1)
  tb(s,f'第{i+1}帧',x,4.76,2.26,.44,20,False,GR,PP_ALIGN.CENTER)
caption(s,'帧率25 fps：每秒显示25张画面。')

s=slide('10分钟、每秒25帧，要保存多少张画面？','从学生已经理解的“视频是连续图像”计算总帧数。先让学生口算25×60×10=15000帧。',2)
tb(s,'25帧/秒 × 60秒/分钟 × 10分钟',.84,2.24,11.8,.85,27,True)
tb(s,'= 15,000 帧',.92,3.57,11.5,1.0,42,True,A)
caption(s,'无压缩视频 = 单帧图像数据量 × 总帧数。',5.37)

q4='10分钟的无声视频：480×320像素，24 bit/像素，\n帧率25 fps。假设逐帧未压缩存储，\n总数据量约多少GiB？'
question(4,q4,'先算每帧，再算总帧数。','来源实录3。不要把帧率误认为采样率，一帧是一张完整图。给3分钟列式，计算可用计算器。',2)
reveal(4,q4,'答案：约6.44 GiB。','480×320×24÷8×25×600 = 6,912,000,000 B；÷1024³≈6.44 GiB。','把公式拆成两段并让学生自己解释。按十进制GB约6.912 GB。',2)

s=slide('可是我们看的10分钟视频，为什么远没有6.44 GiB？','核心认知冲突，不要预先写“因为压缩”。问：是不是算错？为什么在线短视频能很快播放？给学生发言时间。',2)
tb(s,'刚才的计算错了吗？',.88,2.07,11.8,1.0,34,True,A)
tb(s,'如果公式没错，实际文件为什么可能小得多？',.90,3.61,11.7,1.70,28)
caption(s,'请提出至少一种可以检验的解释。')

s=slide('因为我们算的是“逐帧未压缩数据量”。','强调不是估计所有真实视频的固定大小。实际视频通常通过压缩编码，如利用相邻帧相似部分与人眼不敏感的某些信息。高中只解释为什么可能缩小。',2)
diagram_label(s,'未压缩图像帧',.91,2.48,5.02,A)
step_arrow(s,6.03,2.85,7.11,2.85)
diagram_label(s,'重新编码后的视频',7.19,2.48,5.1,C)
tb(s,'每一帧都要保存完整的原始像素数据',.94,4.06,5.03,1.18,21)
tb(s,'省去大量重复表示，数据可能显著减小',7.22,4.06,5.12,1.18,21)
caption(s,'数据量计算仍正确，它描述的是题设中的“未压缩”情况。')

s=slide('为什么要压缩？先想两种真实需求。','实录3行李箱类比：整理行李节约空间；通过网络分享文件缩短传输时间。学生可以自举例，不要说压缩必然提升感知画质。',2)
compare(s,'节省存储空间','同样容量的手机或硬盘，可以保存更多数据。','方便网络传输','需要传输的数据更少，通常能更快下载或播放。','压缩的两个目的：节省空间、提高传输效率。')

s=slide('“中国”重复100次，有没有更简短的写法？','实录3的经典冗余案例。向学生展示短短的“（中国）×100”与重复文字。说明这仅是启发性的编码例子，不代表真实RAR算法。',2)
tb(s,'中国中国中国中国中国中国……',.86,2.16,11.8,1.18,30,True)
step_arrow(s,3.60,4.03,7.05,4.03)
tb(s,'“中国” × 100',7.44,3.62,4.70,1.12,33,True,A)
caption(s,'在不改变内容的前提下，用规则和次数来表示重复部分。')

s=slide('为什么重复内容更容易压缩？','把“冗余”与重复的内容联系：A有规律，B无规律，但原长度相同。实际压缩算法实现不同，因此“通常更易”而非必然。',2)
compare(s,'高度重复的内容','AAAAAAAAAA…\n\n容易用简短规则表示。','变化复杂的内容','AX3Q9P8F2…\n\n难以找到同样简短的规则。','冗余越多，使用合适的算法时，通常越容易压缩。')

s=slide('预测：两份都是1000 Byte，压缩后谁更小？','可以用Python标准库zlib演示，代码放Notes。生成1000个A与1000个伪随机字节（固定seed），比较同一无损压缩方法所得字节数并检验完全还原。先预测再运行。',2)
tb(s,'文件A：1000个“A”',.87,2.19,5.34,.66,27,True,A)
tb(s,'文件B：1000个随机字节',7.02,2.19,5.28,.68,26,True,C)
tb(s,'原始大小相同，都等于1000 Byte。',.88,3.65,11.5,.66,24)
caption(s,'请预测：压缩后 A、B 谁更小？能够完整恢复吗？',5.54)

rng=random.Random(2026)
randbytes=bytes(rng.randrange(256) for _ in range(1000))
size_a=len(zlib.compress(b'A'*1000));size_b=len(zlib.compress(randbytes))
s=slide('实验验证：为什么同样大小的文件效果不同？',f'演示代码（Windows课堂pt环境，Python标准库即可）：\nimport zlib, random\na=b"A"*1000\nr=random.Random(2026)\nb=bytes(r.randrange(256) for _ in range(1000))\nfor name,x in [("A",a),("B",b)]:\n    packed=zlib.compress(x)\n    print(name,len(x),len(packed),zlib.decompress(packed)==x)\n本次固定随机种子得到A={size_a} Byte，B={size_b} Byte。学生关注重复和冗余；算法细节不展开。',2)
compare(s,'A：重复的内容',f'原始大小：1000 Byte\n压缩后：{size_a} Byte','B：复杂的内容',f'原始大小：1000 Byte\n压缩后：{size_b} Byte','两者解压后都能一字节不差地还原：这就是无损压缩。')

s=slide('压缩到底改变了什么？','实录3解释重新编码：数据压缩主要是改变表示方式，无损靠高效编码，允许有损时还可舍弃部分信息。目的与实质分开讲。',2)
tb(s,'数据本身在计算机中仍是0和1。',.89,2.01,11.7,.8,29,True)
tb(s,'压缩：用更少的比特，重新表示相同或近似相同的内容。',.91,3.28,11.7,1.33,27,True,A)
caption(s,'方法有很多，今天重点研究能不能完整还原。')

s=slide('两种压缩：怎样区分“无损”和“有损”？','先提出是否可以还原原始信息，再定义无损有损。注意有损不一定人类完全无感，有时肉眼肉耳可觉察明显差异。',2)
compare(s,'无损压缩','还原以后，与原始数据完全一致。\n\n适合：文档、程序、精确备份。','有损压缩','还原以后，无法与原始数据完全一致。\n\n适合：可接受部分失真的音视频。','判断标准：解码/解压后，能否完整得到原始数据？')

s=slide('ZIP、MP3、WAV、MP4，它们属于同一类吗？','从实录拓展一点格式分层：ZIP/RAR是典型无损归档，MP3是常见有损音频编码，WAV是音频容器可装PCM，MP4是多媒体容器可含有损或无损编码。不要说全部文件扩展名都代表一种压缩算法。',2)
for i,(a,b,col) in enumerate((('ZIP / RAR','通常为无损归档格式',A),('MP3','常见有损音频编码',C),('WAV','容器；常见未压缩PCM',A),('MP4','容器；常装有损视频',C))):
  x=.90+(i%2)*6.1;y=2.04+(i//2)*1.79
  tb(s,a,x,y,5.30,.56,27,True,col);tb(s,b,x,y+.63,5.33,.75,21)
caption(s,'格式名和压缩方法不总是一回事。')

s=slide('三幅相同大小的图，谁通常最容易无损压缩？','实录3三种图像案例。画同尺寸位图：纯色、规则色块、较复杂花纹（不是声称真实风景照片）。首先预测排序，再归纳冗余。',2)
for group,(lab,t) in enumerate((('A 纯色',0),('B 色块',1),('C 复杂图案',2))):
  x=.93+group*4.20;tb(s,lab,x,1.88,3.8,.48,25,True)
  for iy in range(6):
    for ix in range(9):
      if t==0:co=A
      elif t==1:co=A if (ix//3+iy//3)%2==0 else C
      else:co=A if ((ix*7+iy*11)%9)<4 else C
      box(s,x+ix*.39,2.77+iy*.39,.365,.365,co)
  tb(s,'同一尺寸与位深',x,5.41,3.7,.43,19,False,GR)
caption(s,'想一想：哪幅图最容易用几条规则完整表示？')

q5='A纯色、B规则色块、C复杂花纹，尺寸与位深相同。\n分别用相同无损方法压缩，谁通常最小？为什么？'
question(5,q5,'不仅选择，还要解释“重复”和“规律”。','此题迁移文字实验到图像。先解释冗余，再揭示。',2)
reveal(5,q5,'答案：A通常最小。','纯色图像的重复与规律性最强，一般更容易用简短规则表示。','强调不是一切文件和算法都绝对如此。',2)

s=slide('同一张图：BMP与JPG有什么不同？','实录3用同一花朵比较BMP和JPG。用简洁像素格子示意一张图被近似编码；JPEG不只是“删掉相似像素”，要避免这种错误的字面理解。也不能断言所有BMP必然未压缩。',2)
compare(s,'BMP（课堂例：未压缩）','保存每个像素的颜色信息。\n\n可能占用较多空间。','JPG（常见有损）','对图像进行近似表示。\n\n文件可能明显更小，但不能逐像素完全还原。','有损不等于看起来一定明显变差；质量与参数有关。')

s=slide('已经压成MP3、JPG，再用RAR压一遍呢？','实录3脱水比喻，但不把其当算法。说明高效压缩过的数据通常不容易进一步缩小，可能因为已缺少可利用冗余。再次有损转码可能进一步恶化质量。',2)
compare(s,'第一次压缩','有些重复或不重要的信息被更简洁地表示，文件大幅减小。','再次压缩','可进一步利用的冗余通常较少，RAR再压往往收效不大。','注意：再次用有损编码重新保存，还可能累积质量损失。')

s=slide('是不是压得越小就一定越好？','实录3结尾的情境决策：长期备份要完整可还原；实时通话要平衡延迟和质量。用对比培养信息社会责任意识，不只追求压缩比。',2)
compare(s,'重要文档长期备份','优先考虑信息完整、能够恢复，以及可靠保存。','实时视频通话','优先兼顾流畅度、延迟、声音画面质量和传输量。','选择编码与压缩方式，要看具体任务，而不是只比文件大小。')

q6='下列说法，哪些正确？\n① 增加编辑音轨，未压缩PCM数据量必然翻倍。\n② 未压缩视频逐帧保存，数据量可能很大。\n③ ZIP解压后可以恢复原始文档。'
question(6,q6,'逐条作答，并解释理由。','综合三种理解：音轨声道、原始视频大小、无损还原。',2)
reveal(6,q6,'答案：①错；②对；③对。','①音轨不等于声道；②每帧的数据量相加可能很大；③ZIP为无损归档。','不接受只给三个对错而不给依据的回答。',2)

s=slide('离开课堂前，你能解释三个为什么吗？','出口评价：任选一个问题让学生用自己的话解释。重点检查概念因果：数字化、视频容量、压缩选择。',2)
for i,q in enumerate(('① 声音为什么能变成一组0和1？','② 小视频为什么可能算出6.44 GiB？','③ 有些文件为什么压缩后变化很小？')):
  tb(s,q,.86,2.00+i*1.35,11.9,.85,26,True,A if i%2==0 else C)

# Appendix, not planned in main 45m
s=slide('可选拓展：采样不足时会发生什么？','本页只作为选讲：混叠。同一离散采样点可能对应不同连续波形，所以低采样率不仅“点稀疏”，可能无法区分某些信号。深入公式留到拓展课程，不作为主线。',1)
for j,fr in enumerate((1.5,3.5)):
  y=1.94+j*2.14;axes(s,1.2,y,10.4,1.52,'波形 A' if j==0 else '波形 B')
  wave(s,1.3,y,9.55,1.40,frequency=fr,color=A if j==0 else C,stroke=1.9)
caption(s,'拓展问题：少量离散样本，是否足以唯一确定连续信号？',6.39)

s=slide('可选拓展：编码和压缩不是同一件事。','这一页可留作下一单元。PCM编码是样本的表示方式；压缩编码改变数据表示量。部分WAV容器可以保存其他编码形式，不能仅凭扩展名断言完全未压缩。',2)
compare(s,'声音编码','让连续声音能够用数字样本记录；也包括后续数据表示。','数据压缩','用更少的比特表达相同或近似相同的信息。','数字化、容器封装、压缩，是相关但不同的概念。')

# Enrich speaker notes with usable teaching moves for the core concept-changing pages.
DETAILS={
  2:('大家想一想，我们能把一个汉字或一张图片保存进电脑。那么正在说话的声音是不是也能保存？它与图像的最大区别在哪里？', '有人会说只要“录下来”就行；追问“录下来”在计算机里究竟记录什么？', '收束到核心问题：原来连续的声音如何变成数字。'),
  4:('请先判断：空气中的声音和电脑里的音频文件，是不是同一种东西？举一个理由。', '“电脑播放出来还是声音，所以是模拟量”是可能的误解；分清存储形式和播放后的物理振动。', '声音属于物理模拟信号；数字文件以离散的数据保存音频。'),
  5:('请观察这条流程。麦克风是不是已经把声波变成0和1了？请指出转换前后各是什么。', '学生可能认为麦克风就是“录音机”。追问：麦克风输出的电压能否在每个时刻连续变化？', '麦克风把声压变化变成模拟电信号，数字化还没有完成。'),
  6:('如果麦克风输出还是连续电信号，计算机怎样读懂它？需要什么设备？', '学生可能猜CPU，先认可其计算作用，再追问哪个音频部件把模拟信号变成离散样本。', '声卡中的ADC完成模数转换。无需记电子线路结构。'),
  8:('看到这条上下起伏的曲线，请猜：它是不是声音在空气里沿着波浪形路线前进？', '提醒固定一个位置，横坐标向右代表时间，纵坐标上下代表压力或电信号偏离参考的变化。', '波形是“某处随时间变化”的图，不是空间中的传播路径。'),
  9:('现在只允许保存若干个数字，你会在哪些时刻记录这条曲线？让同桌先各提出一条规则。', '有人会只选最高点和最低点。追问：如果换一段波形，这个规则是否容易统一实现？', '用固定时间间隔连续记录，形成采样时刻。'),
 10:('我们按固定时间间隔记录了12个点。灰色曲线不再直接保存，每个点有自己的时间和值。', '学生可能把连点线当作新的连续信号。追问：文件实际存的是曲线，还是这些点对应的数字？', '采样的结果是时间上离散的一组样本。'),
 11:('预测：1 Hz每秒几个点？2 Hz每秒几个点？都记录4秒，两个例子的样本数分别是多少？', '注意“赫兹”还可表示声音本身频率；本页说的是“采样的次数”。', '采样率是每秒采样数；更密意味着相邻采样时刻间隔更短。'),
 13:('假设某个时刻取得3.71这样的振幅值。如果设备只准记录8个整数等级，会怎么办？', '学生可能直接回答“四舍五入”，要追问必须先规定可用的代表值与映射规则。', '量化解决幅度的离散化，下一页构建0到7的教学模型。'),
 14:('请同学数一数，从0到7一共几个数字。0算不算一个等级？', '最容易出现7级的误答。让学生逐个数：0、1、2、3、4、5、6、7。', '共8个等级。这里是简化模型，并非真实声压只有非负值。'),
 15:('规定一个简化规则：向下取整。原值3.71会落在哪个等级？屏幕上的红点与蓝点哪个是原值？', '可能误以为蓝点恰好是原始值。追问：3.71与3有什么不同？少了哪些信息？', '量化是用有限代表值近似原值，存在量化误差。'),
 16:('刚才有8级，若只有两个二进制位可以表示多少种不同情况？三个呢？', '让学生亲自枚举00、01、10、11，然后观察000到111共有八种。', '2^3=8，所以这八个等级可以用3 bit表示。'),
 17:('请从第3个样本找起：它的量化值是多少？对应的三位码字是什么？再找第7个。', '注意区分“样本数量10”和“等级数量8”，不能把十个样本说成十个等级。', '编码按时间顺序保存各个样本的定长码字。'),
 19:('请指出采样率与量化位数分别控制这张图的哪个方向。', '如学生认为量化位数决定点的多少，重新回看每一秒有几个采样时刻。', '采样率决定时间方向的密度，量化位数决定幅度方向的精细程度。'),
 23:('复盘完整流程：数字文件里的样本为什么还能被扬声器播放成声音？', '不要讲成简单把点连成折线。输出经过数模转换及适当的信号处理。', '录音是ADC，播放是DAC，音频数字化形成闭环。'),
 25:('请口算每秒2个样本，录4秒，一路声道，共得到多少个样本？每个3位，要几个bit？', '学生可能直接乘完得到24 Byte。请指着八个框数，再从bit换到Byte。', '8个样本、每个3bit，共24bit，等于3Byte。'),
 26:('我们刚才为什么要把四个因素相乘？请逐一说出公式的单位和意义。', '常见错误：1 kHz写成1024 Hz；bit和Byte混用；忘记把分钟换成秒。', '会用公式之前先能解释公式。'),
 34:('再次提问：多加了一条背景音乐音轨，新的未压缩PCM文件一定比原文件大吗？请先投票并说明依据。', '真实课堂里学生曾将音轨误认为声道；暂不判断两种观点谁对。', '下一页让两种解释并存，定位分歧。'),
 35:('有人支持不变，有人支持翻倍。注意我们的目标不是投票决定真理，而是检验关键假设。', '追问：增加的是“编辑音轨”，还是“导出声道”？哪个出现在容量公式里？', '到图示页验证音轨与声道不同。'),
 36:('请指向编辑左边两行，它们是两条音轨；再指向右边的L/R，这才是导出的两路声道。', '如果学生说“左音轨变左声道”，让他想一想人声与伴奏是否都能同时从左右扬声器传出。', '混合不会自动增加声道数，保持四参数则未压缩样本数据量不变。'),
 40:('五张小图中蓝点位置发生变化，连续播放后我们会感到它在运动。你生活里见过翻页动画吗？', '不要把“视觉暂留0.024秒”当成唯一且固定的生理阈值。', '视频可以用连续图像帧与帧率的概念来理解。'),
 42:('先算一帧480×320像素，需要多少bit和Byte；再乘上25帧/秒和600秒。', '学生可能只乘25或只乘10。追问十分钟共有多少秒、多少帧？', '未压缩视频等于单帧数据量乘总帧数。'),
 44:('请你对照刚刚算出的6.44 GiB。我们手机里十分钟的视频，为什么通常远小于它？难道计算公式错了？', '学生可能回答分辨率不同、格式不同、压缩了。记录这几种可能，指出公式中的题设是未压缩逐帧存储。', '这是第二课时核心认知冲突，直到下一页再揭示压缩。'),
 47:('两种写法能不能表示相同文字？“中国”重复一百次与“中国×100”，哪种通常更省符号？', '提醒这只是解释冗余的教学模型，不是真实RAR编码的说明。', '规律、重复让人有机会用更短的规则表达同一内容。'),
 49:('两个数据各1000Byte：一个全是A，一个随机变化。请预测谁压缩后更小，以及能否恢复原数据。', '学生可能认为所有文件都会变小。下一页结果显示随机数据在容器开销下甚至略微变大。', '测试无损压缩的冗余利用，而不是展示具体算法。'),
 50:('运行标准库实验或直接展示结果：重复内容1000Byte压后17Byte，随机内容1000Byte压后1011Byte。为什么会这样？', '别把高熵文件变大误判为程序出错：压缩文件自身有格式开销。检查两者解压完全一致。', '这是无损压缩，且压缩效果取决于可利用的规律与编码方式。'),
 52:('请判断：压缩后解开，能不能100%恢复原始字节？把这个问题作为分类标准。', '有损不等于“损坏文件”；它在可接受的前提下放弃部分原始信息。', '无损可完全恢复，有损则不能恢复到原数据完全一致。'),
 54:('这三幅图显示的是相同大小、相同位深的像素网格，哪张最容易找到重复的规律？', '强调“通常”和“对合适的算法而言”，避免将图像颜色多少当作绝对压缩决定因素。', '纯色数据具有明显规律，适合无损压缩；比较结论再迁移到实际图像。'),
 59:('某个班级要长期备份源代码，另一个班级要用手机实时视频通话。两种场景需要的编码选择一样吗？', '引导学生考虑完整性、质量、延迟与存储，不要只说压缩得越小越好。', '算法与压缩方法要与任务目标匹配。'),
}
for idx,sld in enumerate(p.slides,1):
    original=sld.notes_slide.notes_text_frame.text.strip()
    if idx in DETAILS:
        say,misunderstanding,close=DETAILS[idx]
        sld.notes_slide.notes_text_frame.text=(
           '【设计意图】'+original+'\n'
           '【教师可以这样讲】'+say+'\n'
           '【预判回答与追问】'+misunderstanding+'\n'
           '【本页收束】'+close)
    else:
        sld.notes_slide.notes_text_frame.text='【教师讲述／组织】'+original+'\n【反馈】先听学生如何解释，再决定是否揭示结论；注意控制本页停留时间。'

# Metadata/verification
p.core_properties.title='声音编码与数据压缩 v6｜两课时课堂版'
p.core_properties.subject='高中信息科技：声音数字化、音频与视频数据量、无损和有损压缩'
p.core_properties.keywords='采样,量化,编码,PCM,声道,音轨,数据压缩,冗余'
assert len(p.slides)==len(SLIDEDATA)
assert all(s.notes_slide.notes_text_frame.text.strip() for s in p.slides)
OUTPUT.parent.mkdir(parents=True,exist_ok=True)
p.save(str(OUTPUT))
print('BUILT',OUTPUT, 'slides',len(p.slides),'bytes',OUTPUT.stat().st_size,'experiment',size_a,size_b)
print('SECTION counts', {k:sum(1 for _,lesson,_ in SLIDEDATA if lesson==k) for k in (1,2)})
for n,lesson,title in SLIDEDATA:print(f'{n:02d} [{lesson}] {title}')