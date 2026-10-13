#!/usr/bin/env python3
"""Ten sequential rounds: review, improve, audit, commit and push each PPTX."""
from pathlib import Path
import os, json, math, hashlib, io, subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image, ImageDraw

BASE=Path(__file__).resolve().parents[2]
AUD=Path(__file__).resolve().parent
IN=BASE/"1-2-3-audio-encoding-v8.1.pptx"
PREFIX="1-2-3-audio-encoding-v8.3."
V=RGBColor(98,81,177);ACC=RGBColor(140,100,225);CYAN=RGBColor(0,176,240)
DARK=RGBColor(38,38,38);MUTE=RGBColor(113,113,113);LIGHT=RGBColor(226,222,242)
WHITE=RGBColor(255,255,255);SOFT=RGBColor(248,246,254);RED=RGBColor(255,0,0)
REG='Alibaba PuHuiTi 3.0 55 Regular';BOLD='Alibaba PuHuiTi 3.0 115 Black'
ISSUES={
1:("P1","模拟量与数字量主要靠文字区分","第4页以连续波形、离散样本同轴比较替换文字"),
2:("P1","麦克风输出连续信号的过程过于抽象","第5页重绘声源、麦克风与连续电信号"),
3:("P1","录音和播放两条转换路径辨识度不足","第7页以紫/青双色双行流程图重绘"),
4:("P1","采样率与量化位数的二维差异不够醒目","第24页将横向样本点与纵向量化等级并列"),
5:("P0","3.71量化示例缺乏清晰的分段归属证据","第17页用0-8数轴和精确落点呈现[3,4)区间"),
6:("P1","PCM容量公式没有显示乘数的具体单位意义","第31页重新绘制四因素对应关系"),
7:("P1","音轨和声道只以标签区分，缺少信号图","第41页以两条音轨和两条输出声道可视对照"),
8:("P0","连续视频帧的小球位置变化不易在投影上分辨","第45页以固定背景、清晰位移重绘五帧序列"),
9:("P1","BMP/JPG只用文字对照，未呈现可观察的有损证据","第63页新增同源原图与低质量JPEG的实际对照"),
10:("P1","三类图像数据冗余度的视觉比较过小","第60页重绘大幅同尺寸纯色/规律色块/复杂图案，复核答案颜色")
}
def pt(v):return Inches(v)
def txt(s,t,x,y,w,h,n=22,b=False,col=DARK,align=None):
    sh=s.shapes.add_textbox(pt(x),pt(y),pt(w),pt(h))
    f=sh.text_frame;f.clear();f.word_wrap=True;f.vertical_anchor=MSO_ANCHOR.MIDDLE
    f.margin_left=f.margin_right=Pt(0);f.margin_top=f.margin_bottom=Pt(0)
    for i,v in enumerate(t.split("\n")):
        p=f.paragraphs[0] if i==0 else f.add_paragraph()
        p.text=v;p.font.name=BOLD if b else REG;p.font.size=Pt(n)
        p.font.bold=b;p.font.color.rgb=col;p.space_after=Pt(0)
        if align is not None:p.alignment=align
    return sh
def rect(s,x,y,w,h,fill=None,stroke=None,width=1.4):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,pt(x),pt(y),pt(w),pt(h))
    if fill:sh.fill.solid();sh.fill.fore_color.rgb=fill
    else:sh.fill.background()
    if stroke:sh.line.color.rgb=stroke;sh.line.width=Pt(width)
    else:sh.line.fill.background()
    return sh
def circ(s,x,y,r,c):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,pt(x-r),pt(y-r),pt(2*r),pt(2*r))
    sh.fill.solid();sh.fill.fore_color.rgb=c;sh.line.fill.background()
    return sh
def seg(s,x1,y1,x2,y2,c=ACC,w=2):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,pt(x1),pt(y1),pt(x2),pt(y2))
    sh.line.color.rgb=c;sh.line.width=Pt(w);return sh
def wave(s,x,y,w,h,cycles=1.2,phase=0,c=MUTE,n=85):
    for j in range(n):
        t=j/n;u=(j+1)/n
        seg(s,x+t*w,y+h*(.5-.38*math.sin(2*math.pi*(cycles*t)+phase)),
             x+u*w,y+h*(.5-.38*math.sin(2*math.pi*(cycles*u)+phase)),c,2)
def clear(s):
    # Keep only title and master-frame rails, replace previous lesson content.
    for sh in list(s.shapes):
        yy=sh.top/914400
        if 1.55<=yy<7.1:
            node=sh._element;node.getparent().remove(node)
def footer(s,value):txt(s,value,.86,6.40,11.70,.42,17,False,ACC)
def note(s,value):
    f=s.notes_slide.notes_text_frame;f.text=(f.text or '').strip()+"\n\n[视觉优化] "+value
def addimg(s,img,x,y,w,h):
    b=io.BytesIO();img.save(b,format='PNG');b.seek(0)
    s.shapes.add_picture(b,pt(x),pt(y),width=pt(w),height=pt(h))
def source_image():
    im=Image.new('RGB',(640,350),(239,246,251));d=ImageDraw.Draw(im)
    d.rectangle((0,180,640,350),fill=(71,130,105))
    d.polygon([(0,178),(112,100),(200,154),(292,65),(422,180)],fill=(79,113,165))
    d.polygon([(180,180),(335,75),(492,181)],fill=(55,90,120))
    d.ellipse((480,35,572,127),fill=(246,206,72))
    for j in range(46):
        x=j*17%640;y=180+j*53%170
        c=(56+j%7*9,105+j%5*8,65+j%8*7)
        d.line((x,y,x+11,y-31),fill=c,width=3)
    for j in range(22):
        y=202+j*6;d.line((0,y,640,y),fill=(64+j%4*3,135+j%3*3,110),width=1)
    for i in range(50):
        x=i*14;d.ellipse((x%640,155+(i*13)%145,x%640+4,160+(i*13)%145),fill=(243,226,132))
    return im
def edit(p,k):
    sl=p.slides
    if k==1:
        s=sl[3];clear(s)
        txt(s,'现实声音：连续变化',.88,1.84,5.4,.55,25,True,ACC)
        txt(s,'电脑保存：离散数字',7,1.84,5.4,.55,25,True,CYAN)
        for x in (.95,7.05):
            seg(s,x,4.77,x+5.2,4.77,MUTE,1.5);seg(s,x,2.59,x,4.90,MUTE,1.3)
        wave(s,1.05,2.83,5.08,1.67,1.65,c=V)
        wave(s,7.15,2.83,5.07,1.67,1.65,c=RGBColor(208,208,213))
        for i in range(12):
            x=7.18+i*5/11;y=3.665-1.67*.38*math.sin(2*math.pi*(i/11*1.65))
            seg(s,x,4.77,x,y,CYAN,1.2);circ(s,x,y,.067,CYAN)
        txt(s,'同一位置的信号随时间变化',.95,5.25,5.4,.55,20)
        txt(s,'只记录特定时刻的一组数值',7.03,5.25,5.3,.55,20)
        footer(s,'连续信号 → 等间隔采样 → 离散样本；图示是同一段信号。')
        note(s,'先让学生观察两图，自己说出波形与离散记录的区别。')
    if k==2:
        s=sl[4];clear(s)
        txt(s,'声源振动',.85,1.86,3.65,.53,25,True,ACC)
        txt(s,'麦克风感受声压',4.91,1.86,3.65,.53,25,True)
        txt(s,'输出连续电信号',9.02,1.86,3.55,.53,25,True,CYAN)
        for i in range(5):
            x=1.15+i*.26;seg(s,x,3.10+(i%2)*.16,x,4.25-(i%2)*.16,ACC,3)
        for j in range(3):
            x=2.65+j*.45
            seg(s,x,3.04,x+.16,3.40,ACC,2);seg(s,x+.16,3.40,x,3.76,ACC,2)
        seg(s,4.25,3.55,4.78,3.55,ACC,3)
        circ(s,6.47,3.63,.52,RGBColor(233,231,239))
        seg(s,6.47,4.15,6.47,4.58,MUTE,3)
        seg(s,7.25,3.55,8.71,3.55,ACC,2.5)
        seg(s,9.01,4.55,12.18,4.55,MUTE,1.2)
        wave(s,9.02,2.72,3.16,1.82,1.35,c=CYAN)
        txt(s,'空气中的变化',.98,5.23,3.3,.53,19)
        txt(s,'转换为电信号',5.12,5.23,3.1,.53,19)
        txt(s,'仍是连续变化的',9.14,5.23,3.1,.53,19)
        footer(s,'麦克风完成“声压变化 → 电信号”；将模拟信号数字化由ADC完成。')
        note(s,'波阵面而非传播路径。麦克风不是二进制编码器。')
    if k==3:
        s=sl[6];clear(s)
        txt(s,'录音：模拟 → 数字',.88,1.69,5,.55,25,True,ACC)
        txt(s,'播放：数字 → 模拟',.88,4.18,5,.55,25,True,CYAN)
        for y,labels,color in ((2.46,['声波','麦克风','ADC','数字文件'],ACC),
                               (4.92,['数字文件','DAC','扬声器','声音'],CYAN)):
            for i,label in enumerate(labels):
                x=.93+3.08*i
                rect(s,x,y,2.15,.71,SOFT if color==ACC else WHITE,color)
                txt(s,label,x,y+.13,2.15,.42,20,True,color,PP_ALIGN.CENTER)
                if i<3:seg(s,x+2.19,y+.35,x+3.0,y+.35,color,2.2)
        footer(s,'思考：麦克风和扬声器、ADC和DAC，各负责什么转换？')
        note(s,'两条平行流程对应录音与播放。')
    if k==4:
        s=sl[23];clear(s)
        txt(s,'采样率：横向有多少次记录？',.85,1.74,5.6,.55,24,True,ACC)
        txt(s,'量化位数：纵向有多少等级？',6.88,1.74,5.8,.55,24,True,CYAN)
        for x in (.88,7.0):
            seg(s,x+.14,4.70,x+5.12,4.70,MUTE,1.3)
            seg(s,x+.14,2.45,x+.14,4.85,MUTE,1.3)
            wave(s,x+.18,2.64,4.90,1.73,1.35,c=MUTE)
        for j in range(11):
            x=1.06+j*.49;y=3.505-1.73*.38*math.sin(2*math.pi*j/10*1.35)
            circ(s,x,y,.07,ACC)
        for j in range(8):
            y=2.55+j*.27;seg(s,7.15,y,12.12,y,CYAN,1.5)
        txt(s,'时间位置变密：样本个数变化',.95,5.32,5.45,.78,21,True,ACC)
        txt(s,'代表值变密：每点存储位数变化',7.03,5.32,5.28,.78,21,True,CYAN)
        footer(s,'重要区别：位数变大，不等于同一时间里采样次数变多。')
        note(s,'看横向点与纵向线：要求学生分别解释对应的参数。')
    if k==5:
        s=sl[16];clear(s)
        txt(s,'输入范围 [0,8) 分成8段：每一段对应一个等级。',.90,1.80,11.5,.63,23)
        x0=.99;unit=1.39;y=3.18
        for i in range(8):
            rect(s,x0+i*unit,y,unit-.04,.72,RGBColor(239,233,255) if i==3 else WHITE,ACC if i==3 else LIGHT,3 if i==3 else 1.2)
            txt(s,str(i),x0+i*unit,y+.08,unit-.04,.46,23,True,ACC if i==3 else DARK,PP_ALIGN.CENTER)
            txt(s,str(i),x0+i*unit-.08,y+.87,.35,.33,15)
        txt(s,'8',12.02,y+.87,.45,.33,15)
        cx=x0+3.71*unit
        seg(s,cx,2.70,cx,y-.02,RED,3);circ(s,cx,2.68,.075,RED)
        txt(s,'3.71',cx-.46,2.12,1.02,.48,23,True,RED,PP_ALIGN.CENTER)
        txt(s,'红点落在哪一段？请先独立判断等级编号。',.97,4.75,11.2,.74,26,True,ACC)
        footer(s,'先只预测，不公布答案；翻页再验证区间规则。')
        note(s,'让学生先预测，再到下一页验证[3,4)的区间归属。')
    if k==6:
        s=sl[30];clear(s)
        txt(s,'数据量（Byte） = 采样率 × 时长 × 声道数 × 位数 ÷ 8',.90,1.74,11.7,.80,25,True,ACC)
        for i,(a,b) in enumerate([('采样率','每秒几个样本'),('时长','采多少秒'),('声道数','同时记录几路'),('量化位数','每个样本多少bit')]):
            x=.95+3.17*i
            rect(s,x,2.99,2.75,1.54,SOFT,ACC)
            txt(s,a,x+.13,3.17,2.5,.43,23,True,ACC,PP_ALIGN.CENTER)
            txt(s,b,x+.09,3.74,2.58,.49,16,False,DARK,PP_ALIGN.CENTER)
            if i<3:txt(s,'×',x+2.79,3.36,.3,.43,21,True,ACC,PP_ALIGN.CENTER)
        seg(s,1.02,5.03,12.04,5.03,LIGHT,1.5)
        txt(s,'前3项相乘 → 样本总数；再乘bit/样本 → bit；最后÷8 → Byte。',.94,5.38,11.65,.64,20,True)
        footer(s,'单位检验：Hz×s×声道×bit÷8，最终为什么得到Byte？')
        note(s,'从样本总数开始推导数据量，解释每个乘数的单位。')
    if k==7:
        s=sl[40];clear(s)
        txt(s,'编辑工程：两条音轨',.87,1.82,5.2,.60,25,True,ACC)
        txt(s,'混音导出：两个声道',7.33,1.82,5.2,.60,25,True,CYAN)
        for i,(label,color) in enumerate([('音轨1：诗朗诵',ACC),('音轨2：背景音乐',CYAN)]):
            y=2.72+i*1.11
            rect(s,.94,y,4.9,.86,WHITE,color,2)
            txt(s,label,1.06,y+.16,2.6,.36,16,True,color)
            wave(s,3.63,y+.10,2.03,.61,2.4 if i==0 else 1.0,phase=.6*i,c=color)
        seg(s,6.02,3.52,7.05,3.52,ACC,3);txt(s,'混合',6.05,3.02,.98,.43,18,True,ACC,PP_ALIGN.CENTER)
        for i,(label,color) in enumerate([('左声道 L',CYAN),('右声道 R',CYAN)]):
            y=2.72+i*1.11
            rect(s,7.26,y,5.0,.86,WHITE,color,2)
            txt(s,label,7.40,y+.16,1.53,.36,16,True,color)
            wave(s,9.09,y+.10,3.02,.61,1.8,phase=i*.5,c=color)
        footer(s,'两条编辑音轨可以混合成两路输出；决定PCM数据量的是输出声道数。')
        note(s,'波形为区别音轨与输出通道的示意图，不是混音运算曲线。')
    if k==8:
        s=sl[44];clear(s)
        txt(s,'同一个球，在5个连续时刻的位置变化',.88,1.76,11.56,.57,24,True,ACC)
        for i in range(5):
            x=.80+2.45*i
            rect(s,x,2.55,2.05,2.27,WHITE,LIGHT,1.5)
            seg(s,x+.15,4.27,x+1.90,4.27,LIGHT,1.3)
            circ(s,x+.30+i*.34,3.67,.18,CYAN)
            txt(s,'第'+str(i+1)+'帧',x,4.95,2.05,.45,19,True,DARK,PP_ALIGN.CENTER)
            if i<4:txt(s,'→',x+2.04,3.45,.39,.48,22,True,ACC,PP_ALIGN.CENTER)
        txt(s,'位置逐帧变化（背景和球的大小不变）',.92,5.65,11.3,.50,20,True,ACC)
        footer(s,'25 fps = 每秒显示25帧；“帧”不是声道，也不是采样位数。')
        note(s,'控制变量只有球的位置；动画效果来自连续帧的顺序。')
    if k==9:
        s=sl[62];clear(s)
        im=source_image();b=io.BytesIO();im.save(b,format='JPEG',quality=3,subsampling=2,optimize=False)
        rec=Image.open(io.BytesIO(b.getvalue())).convert('RGB')
        addimg(s,im,.88,2.13,5.52,3.02);addimg(s,rec,6.94,2.13,5.52,3.02)
        txt(s,'同源原图（无损保存示意）',.91,5.35,5.45,.54,20,True,ACC)
        txt(s,'同一图像：低质量JPEG',6.99,5.35,5.42,.54,20,True,CYAN)
        footer(s,'两图尺寸相同；右图仅用于观察明显压缩伪影，不代表所有JPEG质量。')
        note(s,'同一640×350 RGB程序生成图像以quality=3编码JPEG；有损差异随参数改变。')
    if k==10:
        s=sl[59];clear(s)
        txt(s,'三幅图具有相同的宽、高与位深。',.88,1.73,11.7,.58,24,True)
        for i,(lab,col) in enumerate([('A：纯色',ACC),('B：规律色块',CYAN),('C：复杂变化',V)]):
            x=.86+4.06*i;y=2.38;w=3.62;h=2.94
            if i==0:rect(s,x,y,w,h,RGBColor(83,153,196))
            if i==1:
                for j in range(8):
                    for k2 in range(6):
                        rect(s,x+j*w/8,y+k2*h/6,w/8+.004,h/6+.004,
                             RGBColor(76,157,196) if (j//2+k2//2)%2 else RGBColor(186,220,236))
            if i==2:
                for j in range(24):
                    for k2 in range(18):
                        z=(j*131+k2*251+j*k2*97)%255
                        rect(s,x+j*w/24,y+k2*h/18,w/24+.008,h/18+.008,
                             RGBColor(47+z//3,67+(z*3)%145,85+(z*7)%125))
            txt(s,lab,x,5.45,w,.48,21,True,col,PP_ALIGN.CENTER)
        footer(s,'如果只允许无损压缩，哪张图最容易用更短的规则完整表示？')
        note(s,'相同尺寸和色彩模式；内容结构不同。冗余与压缩算法共同决定结果。')
        for n in [33,38,42,48,62,67]:
            ss=sl[n-1]
            for sh in ss.shapes:
                if sh.has_text_frame and sh.text.strip().startswith('答案：'):
                    for para in sh.text_frame.paragraphs:
                        for run in para.runs:
                            run.font.bold=True;run.font.color.rgb=ACC
def audit(p,k,before,after):
    assert len(p.slides)==70
    titles=[next((sh.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text.strip()),'') for s in p.slides]
    for i in range(1,7):assert titles.count('练习题 '+str(i))==2
    assert all(s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip() for s in p.slides)
    from zipfile import ZipFile
    with ZipFile(after) as z: assert z.testzip() is None
    checksum=lambda f:hashlib.sha256(f.read_bytes()).hexdigest()
    j={'round':k,'severity':ISSUES[k][0],'identified_issue':ISSUES[k][1],'refinement':ISSUES[k][2],
       'slides':70,'notes':70,'question_answer_pairs':6,
       'input_sha256':checksum(before),'output_sha256':checksum(after),
       'method':'PowerPoint native diagram; round 9 actual same-source PNG/JPEG image pair'}
    assert j['input_sha256']!=j['output_sha256']
    (AUD/f'round-{k:02d}-audit.json').write_text(json.dumps(j,ensure_ascii=False,indent=2),encoding='utf-8')
    return j
def main():
    assert IN.exists(),IN
    old=IN
    do_push=os.getenv('RUN_PUSH')=='1'
    for k in range(1,11):
        p=Presentation(old);assert len(p.slides)==70
        edit(p,k)
        p.core_properties.title='声音编码与数据压缩·视觉迭代v8.3.'+str(k)
        path=BASE/(PREFIX+str(k)+'.pptx')
        p.save(path)
        a=audit(p,k,old,path)
        print('ROUND VERIFIED',k,a['severity'],a['refinement'],path.stat().st_size,flush=True)
        if do_push:
            subprocess.run(['git','add',str(path.relative_to(Path.cwd())),
                           str((AUD/f'round-{k:02d}-audit.json').relative_to(Path.cwd()))],check=True)
            subprocess.run(['git','commit','-m',f'visual(audio): finish round {k:02d} and publish '+path.name],check=True)
            subprocess.run(['git','push','origin','HEAD:main'],check=True)
        old=path
    print('ALL TEN ROUNDS COMPLETE',old,flush=True)
if __name__=='__main__':main()
