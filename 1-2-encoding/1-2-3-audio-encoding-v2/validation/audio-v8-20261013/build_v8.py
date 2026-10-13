#!/usr/bin/env python3
"""Rebuild audio v8 from the committed, authentic v7, in 3 audited passes."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import json,zipfile,hashlib

ROOT=Path(__file__).resolve().parents[2]
HERE=Path(__file__).resolve().parent
src=ROOT/"1-2-3-audio-encoding-v7.pptx"
ppt=Presentation(src)
assert len(ppt.slides)==67, f"Need 67 authentic v7 slides, got {len(ppt.slides)}"
ids=[x.id for x in ppt.slides._sldIdLst]
V=RGBColor(98,81,177);ACC=RGBColor(140,100,225)
INK=RGBColor(38,38,38);CYAN=RGBColor(0,176,240)
WHITE=RGBColor(255,255,255);GRAY=RGBColor(130,130,130)
LIGHT=RGBColor(221,221,231);SOFT=RGBColor(243,239,255)
REG="Alibaba PuHuiTi 3.0 55 Regular"
BLACK="Alibaba PuHuiTi 3.0 115 Black"

def orig(n):
    return next(sl for sl in ppt.slides if sl.slide_id==ids[n-1])

def textbox(sl,value,x,y,w,h,size=23,bold=False,color=INK,align=None):
    s=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    f=s.text_frame;f.clear();f.word_wrap=True;f.margin_left=f.margin_right=Pt(0)
    f.margin_top=f.margin_bottom=Pt(0);f.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=f.paragraphs[0];p.text=value;p.font.name=BLACK if bold else REG
    p.font.bold=bold;p.font.size=Pt(size);p.font.color.rgb=color
    if align:p.alignment=align
    return s

def box(sl,x,y,w,h,border=ACC,fill=None):
    s=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill:s.fill.solid();s.fill.fore_color.rgb=fill
    else:s.fill.background()
    if border:s.line.color.rgb=border;s.line.width=Pt(1.5)
    else:s.line.fill.background()
    return s

def newline(sl,x1,y1,x2,y2):
    s=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    s.line.color.rgb=ACC;s.line.width=Pt(2)

def make(title,notes):
    sl=ppt.slides.add_slide(ppt.slide_layouts[min(6,len(ppt.slide_layouts)-1)])
    for p in list(sl.placeholders):
        e=p._element;e.getparent().remove(e)
    box(sl,0,.04,13.333,.08,None,V)
    box(sl,0,7.39,13.333,.08,None,V)
    textbox(sl,title,.665,.64,12,.85,35,True)
    sl.notes_slide.notes_text_frame.text=notes
    return sl

def append_note(sl,msg):
    txt=sl.notes_slide.notes_text_frame
    txt.text=txt.text.strip()+"\n\n"+msg

def replace_word(sl,old,new):
    found=False
    for shape in sl.shapes:
        if not shape.has_text_frame:continue
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                if old in r.text:
                    r.text=r.text.replace(old,new);found=True
    assert found,(old,new)

def reorder(order):
    items=ppt.slides._sldIdLst
    mapping={e.id:e for e in list(items)}
    assert len(order)==len(mapping) and set(order)==set(mapping)
    for e in list(items):items.remove(e)
    for sid in order:items.append(mapping[sid])

def check(stage,count):
    assert len(ppt.slides)==count
    assert all(sl.has_notes_slide and len(sl.notes_slide.notes_text_frame.text.strip())>35 for sl in ppt.slides)
    text=lambda s:[z.text.strip() for z in s.shapes if z.has_text_frame and z.text.strip()]
    titles=[text(s)[0] for s in ppt.slides]
    demo_c=next(i for i,t in enumerate(titles) if "演示C" in t)
    code=next(i for i,t in enumerate(titles) if "一个样本需要几个二进制位" in t)
    assert code<demo_c
    assert titles.count("练习题 1")==2
    assert titles.count("练习题 2")==2
    assert titles.count("练习题 3")==2
    assert titles.count("练习题 4")==2
    assert titles.count("练习题 5")==2
    assert titles.count("练习题 6")==2
    if stage>=2:
        assert any("演示D" in t for t in titles)
        ix=next(i for i,t in enumerate(titles) if "演示D" in t)
        assert "预测" in titles[ix-1] and "实验验证" in titles[ix+1]
    if stage==3:
        assert "013" not in titles[19] # no placeholder-code mistake
    file=HERE/("round-"+str(stage)+".pptx")
    ppt.save(file)
    with zipfile.ZipFile(file) as f:assert f.testzip() is None
    report={"stage":stage,"slides":len(ppt.slides),"speaker_notes":sum(s.has_notes_slide for s in ppt.slides),"file_sha256":hashlib.sha256(file.read_bytes()).hexdigest(),"practice_pairs":6,"code_introduced_before_demo_c":True}
    (HERE/("audit-round-"+str(stage)+".json")).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    print("PASS",report)
    return file

# REVIEW 1: P0 unclear quantizer rule, P1 demo precedes binary concept, P2 transition notes.
question=make("3.71 应记为等级3，还是等级4？","先显示0~7八级，学生预测3.71是3或4。暂停50秒让学生说出推理规则。再用下页指定的分区规则解决分歧；勿预先提示答案。")
textbox(question,"幅度范围 [0,8) 划成8段，等级编号为 0～7。",.9,1.86,11.8,.72,25)
for i in range(8):
    x=.95+1.50*i
    box(question,x,3.15,1.20,1.00,ACC if i in (3,4) else LIGHT)
    textbox(question,str(i),x,3.31,1.20,.53,30,True,ACC if i in (3,4) else GRAY,PP_ALIGN.CENTER)
textbox(question,"采样值：3.71",.95,4.64,5.3,.74,29,True)
textbox(question,"你预测：进入3号等级，还是4号等级？",.95,5.61,11.5,.8,27,True,ACC)
mapping=make("3.71 为什么写成 011，而不是 100？","逐步指着四个框引导：输入原值3.71，属于区间[3,4)，所以等级编号是3；把3写成三位等长码字011。反例4.15属于L4，码字100。本例幅度[0,8)，Demo C用[-1,1]，逻辑完全相同。")
flow=[("采样原值","3.71"),("输入范围","[3,4)"),("量化等级","L3"),("3 bit码字","011")]
for i,(a,b) in enumerate(flow):
    x=.82+i*3.15
    box(mapping,x,2.46,2.66,1.79,ACC,SOFT if i==3 else WHITE)
    textbox(mapping,a,x+.10,2.67,2.48,.5,21,False,ACC,PP_ALIGN.CENTER)
    textbox(mapping,b,x+.10,3.26,2.48,.65,31,True,ACC if i==3 else INK,PP_ALIGN.CENTER)
    if i<3:newline(mapping,x+2.72,3.35,x+3.08,3.35)
textbox(mapping,"规则：按输入范围确定等级，再用固定长度的0和1记录等级编号。",.86,5.10,11.65,.8,24,True)
textbox(mapping,"验证：输入值4.15对应的码字是什么？",.87,6.37,11.3,.48,21,False,ACC)
replace_word(orig(17),"记录 3","等级 L3")
textbox(orig(17),"规则：[3,4) → L3；区间代表值可取 3.5",1.05,5.90,11.1,.42,18,True,ACC)
for n,note in [
 (16,"讲授0~7是等级编号；不要把取样数、幅度值、等级编号混为一谈，先让学生预测。"),
 (17,"这里按[n,n+1)归入Ln，3.71属于L3；量化代表值可取3.5。让学生回看先前错误猜测。"),
 (18,"该比较图与Notebook 03采用不同示例幅度范围，但都是先采样、划区间、记录等级；先建立3bit=8级再操作。"),
 (19,"三位二进制可以表达8种组合；先问学生最少需要几bit，再揭示全部码字。"),
 (9,"打开远程JupyterLab Notebook 01，在Windows浏览器播放声音，问红线向右是不是声波向右传播；让学生说横轴是时间。"),
 (12,"打开Notebook 02，把波形从100毫秒放大到4毫秒，点击蓝色样本，读出实际时间和幅度；区分极值概览和真实采样点。"),
 (20,"十个样本的表先只聚焦第一列：样本编号、量化等级、3bit码字逐一对应，再请学生解释其它样本。"),
]:
    append_note(orig(n),note)
textbox(orig(43),"球在画面中的位置逐帧变化",2.50,5.53,8.55,.64,22,True,ACC,PP_ALIGN.CENTER)
append_note(orig(43),"五帧是同一个运动小球在五个时刻的状态。帧率25fps是每秒25帧，不应解释为人眼严格生理阈值。")
order=ids[:16]+[question.slide_id]+ids[16:17]+ids[18:19]+[mapping.slide_id,ids[17]]+ids[19:]
reorder(order)
assert len(ppt.slides)==69
check(1,69)

# REVIEW 2: P1 no compression experiment bridge, P1 empirical results lack provenance, P2 classification caveats.
d=make("演示D：同样1000 Byte，压缩后真的不同吗？","运行remote-demos-v7/notebooks/04_压缩前先预测.ipynb（Mac计算zlib，教室Windows浏览器呈现结果）。学生先猜，再运行，再对照下一页17Byte/1011Byte。B可能变大，包含必要编码开销。zlib不是RAR算法本身。")
textbox(d,"同样是 1000 Byte：",.90,1.78,11.5,.7,26,True)
for i,(name,desc) in enumerate([("文件 A","1000个重复的字节“A”"),("文件 B","1000个随机变化的字节")]):
    x=.9+6.35*i
    box(d,x,2.83,5.58,1.38,ACC if i==0 else CYAN)
    textbox(d,name,x+.21,3.06,1.48,.56,23,True,ACC if i==0 else CYAN)
    textbox(d,desc,x+1.66,3.06,3.58,.57,19)
textbox(d,"① 先预测 → ② 运行Notebook的zlib压缩 → ③ 比较结果",.88,4.72,11.62,.93,24,True,ACC)
textbox(d,"可追问：压缩后的文件有没有可能比原来更大？",.88,6.08,11.5,.58,21)
order=[s.slide_id for s in ppt.slides]
idx=order.index(ids[51])
order.remove(d.slide_id);order.insert(idx+1,d.slide_id);reorder(order)
append_note(orig(53),"17Byte与1011Byte是采用Python标准库zlib、固定随机种子的这一次课堂实验结果。数值取决于内容与算法，不是普遍常数；B数据因格式开销而可能变大。")
append_note(orig(56),"ZIP/RAR属于无损归档，MP3是有损音频编码，WAV和MP4属于容器（MP4常装有损编码视频）。文件后缀与编码算法不是完全同一层级。")
append_note(orig(60),"BMP这里只指未压缩实例，JPG通常有损；文件大小还取决于参数和内容。")
replace_word(orig(48),"省去大量重复表示，数据可能显著减小","利用帧间相似和更有效的编码，数据可能显著减小")
append_note(orig(48),"实际压缩技术还包括多种预测与表示方式；本课只要求理解重复规律允许节省比特，不学习具体codec算法。")
check(2,70)

# REVIEW 3: Speaker notes focus on what students predict, how evidence changes their minds, and what counts as understanding.
notes={
 2:"先请学生复述文字和图片怎样编码，再让学生预测麦克风后是什么、计算机需要什么。教师不预告完整三步。",
 4:"对比模拟声波与文件中的数字，问连续变化的声音如何变成离散数据；先让学生说两者有何不同。",
 5:"麦克风负责把声压变化变成连续电信号；先预测“是数字还是模拟？”学生说出依据。",
 6:"声卡ADC把模拟信号变为数字数据；不要求硬件电路细节，用前后变化解释。",
 7:"按录音顺序走一遍，再反向说播放顺序，确认DAC与ADC的作用不同。",
 10:"请学生先在波形上建议记录时刻，再比较哪种规则稳定；引出等时间间隔采样。",
 11:"先观察时间刻度和蓝色点，教师追问记录的是什么量，说明在指定时刻记录幅度。",
 13:"1Hz=每秒1个记录点；2Hz=每秒2个。这里用于放大理解，绝非实际音频录制建议。",
 14:"上下保持同一条曲线不变，只改变时间方向的样本密度，让学生辨认控制变量。",
 15:"讨论3.71这样的幅度怎样用有限个等级近似记录，不让量化等同四舍五入。",
 21:"对比3bit与4bit的等级数，先让学生预测会不会增加样本点个数；解释是纵向可选值更密。",
 22:"分别指横轴时间密度与纵轴量化等级，要求学生指出哪一项由采样率/位数控制。",
 23:"纠正“设置越高听觉效果必然越好”的绝对化说法，讨论保真能力与具体条件。",
 25:"区分数字样本和WAV容器头的作用；WAV不等同于一种有损压缩方法。",
 26:"沿数字样本→DAC→扬声器→声音完成录放闭环，不在主线要求复述滤波细节。",
 27:"先让学生自己说PCM数据量四个因素，而不是直接背公式。",
 28:"让学生数8个样本，每个3bit共24bit，除8得到3Byte，单位帮助理解公式。",
 29:"为什么结果是Byte要÷8？把四因素分别对应实际样本数量与每样本bit数。",
 30:"独立列式后再揭示；提问哪一个乘数是双声道、为什么还要除以8。",
 31:"答案160000 Byte。让学生先说思路，再对照上一页原位置的题干与下方答案。",
 34:"先只看48000、16、2、60，要求学生用自己的话说明每个参数是什么。",
 35:"先从PCM字节数算起，再转换MiB；强调1024²与1000²换算不同。",
 36:"约10.99 MiB；验证11520000 B/1048576。不把带少量文件头的WAV大小当作PCM样本大小。",
 37:"投票多一条编辑音轨容量是否翻倍；不同学生给出理由，暂不揭示正确答案。",
 38:"两种观点的矛盾来自把编辑音轨当作输出声道。追问最后导出文件是几声道。",
 39:"先看两条轨，再揭示混音之后仍可输出L/R两路数据。",
 40:"回到初始投票，解释四个PCM参数保持不变，因此样本数据量不变。",
 41:"迁移：把立体声变成单声道，其他参数相同，PCM样本数据量减半。",
 43:"球在每一帧的位置不同，同样场景取不同时间画面；fps是每秒帧数。",
 45:"10分钟无声逐帧未压缩视频，先算单帧字节，再算15000帧，单位换算为GiB。",
 46:"答案约6.44GiB，检验帧数、bit转Byte以及B转GiB三步。",
 47:"先让学生用自己看视频的经验提出文件比估算小的解释，借助冲突引出编码压缩。",
 50:"中国重复100次只用“字词×次数”表达是思想模型，不代表RAR或zlib唯一工作机制。",
 51:"重复多往往可用更短的规则表示，但实际压缩要看数据及算法。",
 52:"A和B均1000Byte，学生先预测谁更省空间、B是否可能压完更大。",
 54:"归纳数据压缩是更高效的表示；要不要完整恢复决定无损与有损路径。",
 55:"判别方法：是否能够完全恢复原字节，不用文件后缀或大小作为唯一标准。",
 57:"三图尺寸和位深一致，让学生预测哪个更规律，判断无损压缩的可能差异。",
 58:"练习答案不能只选A，要解释重复和规律性带来更多可利用冗余。",
 59:"答案中的“通常”必须保留，不同压缩方案有例外。",
 61:"已经是MP3/JPG，用RAR压缩通常收益有限，再次有损编码还可能降低质量。",
 62:"文档备份要求完整，视频通话权衡时延/带宽/质量，反对越小必然越好。",
 63:"请学生独立判断三句话，指明每句对应的模型或证据。",
 64:"先让学生口头解释，再展示①错②对③对，强化数据量与ZIP的两种知识。",
 65:"退出条检查解释能力而非术语背诵；答不出则回到图示逐步重建。"
}
for n,ex in notes.items():append_note(orig(n),ex)
for num in [31,36,40,46,59,64]:
    for obj in orig(num).shapes:
        if not obj.has_text_frame or not obj.text.strip().startswith("答案："):continue
        for para in obj.text_frame.paragraphs:
            for run in para.runs:
                run.font.bold=True
                run.font.color.rgb=V
                if run.font.size and run.font.size.pt<22:run.font.size=Pt(22)
for n in [66,67]:append_note(orig(n),"可选拓展：不是45分钟主线必讲内容，有时间且学生已完成必要理解才展示。")
ppt.core_properties.title="声音的数字化、音频视频编码与数据压缩 v8｜两课时课堂版"
ppt.core_properties.subject="高中信息科技：声音数字化与数据压缩"
out=check(3,70)
final=ROOT/"1-2-3-audio-encoding-v8.pptx"
final.write_bytes(out.read_bytes())
print("FINAL",final,final.stat().st_size)
