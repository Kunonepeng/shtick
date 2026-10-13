#!/usr/bin/env python3
"""Rebuild v7 classroom deck from v6 with three browser-demonstration anchors.

No old slide is replaced: original 64 v6 slides, their visuals and notes are preserved.
Inserted slides link the learning question to a concrete JupyterLab demo.
"""
from pathlib import Path
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN

ROOT=Path(__file__).resolve().parent
# For GitHub, this script is in validation/audio-v7-20261013/.
if not (ROOT/'1-2-3-audio-encoding-v6.pptx').exists():
    ROOT=Path(__file__).resolve().parents[2]
SOURCE=ROOT/'1-2-3-audio-encoding-v6.pptx'
if not SOURCE.exists() and (Path(__file__).parent/'v6_source.pptx').exists():
    SOURCE=Path(__file__).parent/'v6_source.pptx'  # local development only
if not SOURCE.exists():
    raise FileNotFoundError(f'Cannot find v6 base deck: {SOURCE}')
OUT=(Path(__file__).parent/'1-2-3-audio-encoding-v7.pptx') if (Path(__file__).parent/'v6_source.pptx').exists() else ROOT/'1-2-3-audio-encoding-v7.pptx'
p=Presentation(str(SOURCE))
assert len(p.slides)==64, len(p.slides)
V=RGBColor(98,81,177)
A=RGBColor(140,100,225)
C=RGBColor(0,176,240)
RED=RGBColor(255,0,0)
INK=RGBColor(38,38,38)
GRAY=RGBColor(128,128,128)
L=RGBColor(226,226,231)
WHITE=RGBColor(255,255,255)
REG='Alibaba PuHuiTi 3.0 55 Regular'
HEAVY='Alibaba PuHuiTi 3.0 115 Black'


def text(sl,s,x,y,w,h,pt=22,col=INK,bold=False,align=None):
    z=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    f=z.text_frame;f.clear();f.word_wrap=True
    f.margin_left=f.margin_right=Pt(0);f.margin_top=f.margin_bottom=Pt(0)
    for i,line in enumerate(s.split('\n')):
        para=f.paragraphs[0] if i==0 else f.add_paragraph()
        para.text=line;para.space_after=Pt(3)
        para.font.name=HEAVY if bold else REG;para.font.size=Pt(pt);para.font.bold=bold;para.font.color.rgb=col
        if align is not None:para.alignment=align
    return z

def rect(sl,x,y,w,h,fill=None,border=None,bw=1):
    shape=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill:shape.fill.solid();shape.fill.fore_color.rgb=fill
    else:shape.fill.background()
    if border:shape.line.color.rgb=border;shape.line.width=Pt(bw)
    else:shape.line.fill.background()
    return shape

def line(sl,x,y,X,Y,col=A,bw=2):
    s=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(X),Inches(Y))
    s.line.color.rgb=col;s.line.width=Pt(bw)
    return s

def dot(sl,x,y,r=.07,col=C,border=None):
    z=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-r),Inches(y-r),Inches(r*2),Inches(r*2))
    z.fill.solid();z.fill.fore_color.rgb=col
    z.line.color.rgb=border or col;z.line.width=Pt(1.5)
    return z

def start(title,notes):
    sl=p.slides.add_slide(p.slide_layouts[min(6,len(p.slide_layouts)-1)])
    for ph in list(sl.placeholders):
        item=ph._element;item.getparent().remove(item)
    rect(sl,0,.04,13.333,.08,V);rect(sl,0,7.39,13.333,.08,V)
    text(sl,title,.67,.63,12,.86,34,INK,True)
    sl.notes_slide.notes_text_frame.text=notes
    return sl

def graph(sl,x,y,w,h,dots=0,zoom=False):
    line(sl,x,y+h/2,x+w,y+h/2,GRAY,1.4)
    line(sl,x+.05,y+.1,x+.05,y+h-.05,GRAY,1.1)
    text(sl,'振幅',x,y-.37,1.25,.32,17,GRAY)
    text(sl,'时间',x+w-1,y+h/2+.14,.90,.32,17,GRAY)
    freq=2.0 if not zoom else 1.4
    xy=[]
    for i in range(101):
        t=i/100
        px=x+.13+t*(w-.4)
        py=y+h/2-.34*h*math.sin(2*math.pi*freq*t+.17)
        xy.append((px,py))
    for a,b in zip(xy[:-1],xy[1:]):line(sl,*a,*b,GRAY,1.7)
    if dots:
        for k in range(dots):
            t=k/dots
            xx=x+.13+t*(w-.4)
            yy=y+h/2-.34*h*math.sin(2*math.pi*freq*t+.17)
            dot(sl,xx,yy,.063,C)

# Add contextual teacher notes to existing v6 slides (do not overwrite them).
patch={
    8:'【v7 Demo A，约2分钟】先保留本页“传播路线吗”的问题，让学生预测；运行 notebooks/01_听见声音_看见波形.ipynb，在教室Windows浏览器中点击播放。指向红色时间光标：波形表示时间变化，不是空间传播路线。Mac不会播放声音，声音从Windows浏览器输出。网络异常时直接用本页静态图。',
    10:'【v7 Demo B，约3分钟】运行 notebooks/02_放大波形_看见采样点.ipynb，将观察范围从100毫秒缩至20、4毫秒，直到看见一个个真实样本点。与完整波形极值包络对照，不要把包络竖线叫采样点。先问“看似连续的曲线，计算机究竟保存了什么？”',
    11:'【v7回扣】演示切换到本页后，要求学生解释采样率和每秒样本数之间的关系。Notebook中的4毫秒局部窗口只是观察尺度。',
    13:'【v7 Demo C预告】现在只解决了采样时刻；纵向幅度可能是任意连续值，需要量化，稍后用实验台在同一坐标中比较。',
    19:'【v7 Demo C，约5分钟】运行 notebooks/03_采样_量化_编码实验台.ipynb，按①原波形→②采样→③量化→④编码进行。学生先预测调高采样率会变哪些点，再预测保持采样率改bit数时横向点数会否变化。实验台使用归一化-1到1幅度，量化等级编号仍为0～2^n-1，与本页0～7教学例的单位尺度不同。',
    20:'【v7实验台回扣】在实验台选择3 bit，点击一个点，要求学生报出等级编号和3 bit码字。注意实验台使用[-1,1]输入范围；本页0～7演示是另一种简化尺度，编码原理相同。',
    21:'【v7学习证据】让学生说出：“采样确定何时记录，量化决定用哪个代表值，编码将编号写成固定长度的二进制码字。”如果答不出，回实验台查看同一个样本，避免只背步骤名称。',
    29:'【v7时间控制】本课新增的三个演示分别约2、3、5分钟，应替换原有重复的静态口头描述；并非在原45分钟上增加10分钟。',
    49:'【v7 Demo D】先预测两份原始都是1000 Byte的文件哪一个被压得更小，再运行 notebooks/04_压缩前先预测.ipynb。Python zlib在Mac执行，Windows浏览器只展示结果，不依赖音频或py5。',
    50:'【v7实际实验】Notebook 04使用确定性随机字节与重复字节，打印压缩前后真实长度。如果随机字节压缩后略大，是格式开销而非代码错误。让学生总结“能够用规律简洁表示的信息通常更容易压缩”。'
}
for no,addon in patch.items():
    s=p.slides[no-1]
    old=s.notes_slide.notes_text_frame.text
    s.notes_slide.notes_text_frame.text=old.rstrip()+"\n\n"+addon

# Build and move new slides from back to front to keep insertion positions stable.

def after(no,sl):
    sid=p.slides._sldIdLst[-1]
    p.slides._sldIdLst.remove(sid)
    p.slides._sldIdLst.insert(no,sid)

s=start('演示C：3 bit变4 bit，横向的点会变多吗？',
        '在已理解3bit与8个等级后，先给学生两种预测。运行notebooks/03_采样_量化_编码实验台.ipynb，固定4 kHz，位数3→4。学生观察：采样时刻不变，纵向可选的量化代表值从8增加到16，码字长度相应变化。此图只是静态示意，真实数值以Notebook演示为准。约5分钟，不与第15～17页全部额外累加。')
text(s,'保持采样时刻不变，只改变量化位数。',.91,1.68,11.4,.65,25,A,True)
for j,(lab,levels,col) in enumerate([('3 bit：8个等级',8,A),('4 bit：16个等级',16,C)]):
    y=2.47+j*1.93
    text(s,lab,.91,y+.35,3.00,.49,22,col,True)
    left=4.24;right=12.12;top=y;bottom=y+1.49
    for k in range(levels):
        yy=top+(bottom-top)*k/max(1,levels-1)
        line(s,left,yy,right,yy,col,1.0)
    # Same times AND same original sample values in both panels.
    for k in range(9):
        xx=left+(right-left)*k/8
        amplitude=.60*math.sin(2*math.pi*k/8+.32)
        yy=(top+bottom)/2-amplitude*(bottom-top)/2
        dot(s,xx,yy,.059,GRAY)
text(s,'灰点的时间位置不变；变化的是纵向可选等级。',1.04,6.57,11.4,.42,20,A,True)
after(19,s)

s=start('演示B：把波形放大，能看见一个个采样点吗？',
        '运行notebooks/02_放大波形_看见采样点.ipynb。先保持完整波形，再把局部显示范围切换到20毫秒和4毫秒。左侧是示意概览，右侧是局部实际样本思路；在Notebook里，右侧点取自真实生成的PCM数据。不要说采样点之间连线是扬声器实际波形。约3分钟。')
text(s,'同一段录音：左边看整体，右边看极短的一段。',.88,1.64,11.4,.65,24,A,True)
text(s,'整段波形',.92,2.40,5.4,.49,24,INK,True)
text(s,'局部放大：真实样本点',7.05,2.40,5.6,.49,24,INK,True)
graph(s,.88,3.06,5.50,2.18)
graph(s,7.05,3.06,5.35,2.18,dots=15,zoom=True)
text(s,'先预测：放大到足够短的时间，会出现什么？',.96,5.97,11.5,.67,23,A,True)
after(10,s)

s=start('演示A：听声音，波形上的红线如何移动？',
        '运行notebooks/01_听见声音_看见波形.ipynb。请教师从Windows浏览器直接点击播放；让学生预测红线表示空间位置还是播放时间，然后暂停再播放。观察教室音箱发声，Mac只是运算和提供数据。约2分钟。网络失败时直接使用本页示意和原第8页。')
text(s,'先预测：红色光标向右，表示声音向右传播吗？',.86,1.66,11.7,.70,25,A,True)
graph(s,1.03,2.66,11.18,2.80)
line(s,7.26,2.83,7.26,5.33,RED,3.2)
text(s,'▶ 由教室浏览器播放',1.00,5.83,6.3,.54,23,INK,True)
text(s,'时间继续向右',8.45,5.83,3.4,.54,23,A,True)
after(8,s)

assert len(p.slides)==67,len(p.slides)
assert all(s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip() for s in p.slides)
p.save(str(OUT))
print('wrote', OUT, OUT.stat().st_size, 'bytes, pages',len(p.slides))
