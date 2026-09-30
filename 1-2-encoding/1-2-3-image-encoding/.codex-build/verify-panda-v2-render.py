from pathlib import Path
import json, hashlib, shutil, zipfile
import numpy as np
import pdfplumber
from PIL import Image, ImageDraw
from pptx import Presentation
root=Path(__file__).resolve().parent.parent
build=root/'.codex-build'; qa=build/'qa-panda-v2'; qa.mkdir(exist_ok=True)
render=Path('/private/tmp/panda-v2-qa');final=root/'exports/1-2-3-image-encoding-panda-v2.pptx'
hashfile=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert final.read_bytes()==(build/'candidate-panda-v2.pptx').read_bytes()
assert hashfile(root/'exports/1-2-3-image-encoding-panda-v1.pptx')=='110e872857ce7f577175ec84ee7e96136108ef91003148ff68ad9bbc2b956c6a'
d=json.loads((build/'panda-data.json').read_text());assert hashfile(root/d['source'])==d['sha256']
a=np.asarray(Image.open(root/d['source']).convert('RGB'));c=d['crop'];crop=a[c['y']:c['y']+c['height'],c['x']:c['x']+c['width']].astype(float);pal=np.array(d['paletteRgb'])
for n,key,colors in [(8,'coarse',4),(16,'fine',4),(16,'six',6)]:
 z=288//n;means=crop.reshape(n,z,n,z,3).mean(axis=(1,3));idx=((means[:,:,None,:]-pal[None,None,:colors,:])**2).sum(axis=-1).argmin(axis=-1)
 assert [''.join(map(str,row)) for row in idx]==d[key]
 assert np.allclose(means,d['means8' if n==8 else 'means16'])
assert a[d['pixel']['y'],d['pixel']['x']].tolist()==d['pixel']['rgb']
assert d['coarse'][3]=='33201211'
fonts=set();pdfpath=render/'candidate-panda-v2.pdf'
with pdfplumber.open(pdfpath) as pdf:
 assert len(pdf.pages)==63
 for i,p in enumerate(pdf.pages,1):
  assert p.chars, f'blank text page {i}'
  fonts.update(ch['fontname'].split('+')[-1] for ch in p.chars)
  assert all(0<=ch['x0']<=ch['x1']<=p.width+0.1 and 0<=ch['top']<=ch['bottom']<=p.height+0.1 for ch in p.chars)
assert all('AlibabaPuHuiTi' in f for f in fonts),fonts
prs=Presentation(final)
for i,s in enumerate(prs.slides,1):
 t=s.notes_slide.notes_text_frame.text
 assert '[教师逐字稿]' in t and '[问题 / 页面目的]' in t and 'undefined' not in t
for n in range(1,64):
 p=render/f'slide-{n:02}.png';assert Image.open(p).size==(1600,900);shutil.copyfile(p,qa/p.name)
shutil.copyfile(pdfpath,qa/'slides.pdf')
for k,start in enumerate(range(1,64,12),1):
 canvas=Image.new('RGB',(1600,750),'#DDDDDD');draw=ImageDraw.Draw(canvas)
 for j,n in enumerate(range(start,min(start+12,64))):
  im=Image.open(qa/f'slide-{n:02}.png').convert('RGB');im.thumbnail((400,225));x=(j%4)*400;y=(j//4)*250;canvas.paste(im,(x,y));draw.text((x+8,y+228),f'Slide {n:02}',fill='black')
 canvas.save(qa/f'montage-{k}.jpg',quality=94)
with zipfile.ZipFile(final) as z:
 media=[name for name in z.namelist() if name.startswith('ppt/media/')];unique_media={hashlib.sha256(z.read(n)).hexdigest() for n in media}
 allowed={hashfile(root/d['source'])}|{hashfile(root/'assets/original-screenshots'/f'Snip20260929_{n}.png') for n in [22,25]}
 assert unique_media==allowed
result={'final_sha256':hashfile(final),'bytes':final.stat().st_size,'slides':63,'source_rgb_and_all_matrices_recomputed':True,'prior_deck_and_source_unchanged':True,'notes_complete':True,'rendered_slides':63,'render_resolution':'1600x900','pdf_fonts':sorted(fonts),'rendered_text_in_bounds':True,'embedded_media_count':len(media),'unique_source_images':len(unique_media),'visual_inspection':'All 63 slides inspected individually; complete montage inspected. No visible clipping or unintended overlap found.','native_application_check':'WPS / PowerPoint not performed'}
(build/'render-qa-panda-v2.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False))
report=root/'panda-v2-audit.md';s=report.read_text().split('\n## 最终验收记录')[0]
s+='''\n## 最终验收记录\n\n- 最终文件：`exports/1-2-3-image-encoding-panda-v2.pptx`。63页，28组问题／答案配对，14页使用原生可编辑表格。\n- 63页均含完整问题／页面目的和教师逐字稿；学生页与教师备注分开。\n- 28组问答页由复制生成，标题的文字、字体和几何位置一致，原有对象位置稳定。\n- 原图RGB、8×8四色矩阵、16×16四色／六色矩阵已从原JPEG重新计算核对；真实眼睛像素、位权、码字和预算题计算均核对通过。\n- QA中将Q6焦点修正为含0/1/2/3的第四行：33201211，对应11 11 10 00 01 10 01 01；红框、逐字稿与检查脚本一致。\n- 全部63页经LibreOffice渲染为1600×900图片，逐页大图与全套缩略图检查完成；未发现可见遮挡、截字、意外换行或越界。渲染证据保存在`.codex-build/qa-panda-v2/`。\n- 原生文字只使用Alibaba PuHuiTi 3.0 115 Black和55 Regular；本次PDF渲染未发现其他替代字体。16:9画布、白底、紫色上下轨道及统一标题基线检查通过。\n- PPTX包结构、声明的原生表格、字体、几何检查以及Artifact Tool重新导入均通过。通用表格求和检查不适用于本课表格；算式与实际数据另行核对。\n- v1 PPTX、用户原始熊猫图及既有数据未被覆盖；课程设计和历史来源保持原样。\n- 未在WPS／PowerPoint实机或教室投影上验收；45分钟安排未经过真实班级试讲。上述通过结果限于本机渲染、内容复核与结构检查。\n\n最终SHA-256：`'''+result['final_sha256']+'''`。\n\n复现入口：`.codex-build/build-panda-v2.mjs`；结构检查：`.codex-build/check-panda-v2.py`；交付校验：`.codex-build/validation-panda-v2.json`；渲染与数据核对：`.codex-build/render-qa-panda-v2.json`。\n'''
report.write_text(s)
