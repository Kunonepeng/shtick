#!/usr/bin/env python3
import argparse, json, re, zipfile, posixpath
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET
from difflib import SequenceMatcher

from pptx import Presentation

EMU_PER_INCH = 914400
NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
}

EXPECTED = {
    'width_in': 13.333,
    'height_in': 7.5,
    'frame_violet': '6251B1',
    'rail_h_in': 8/72,
    'title_x': 0.665,
    'title_y': 0.665,
    'title_w': 12.0,
    'title_pt': 36,
    'title_font_fragment': 'Alibaba PuHuiTi 3.0 115 Black',
}

TECH_REQUIRED = [
    ('ASCII 7-bit', ['ASCII', '7 bit']),
    ('GB2312 国标码/机内码', ['国标码', '机内码']),
    ('Unicode code point', ['Unicode', 'U+4F60']),
    ('UTF-8', ['UTF-8']),
    ('UTF-16', ['UTF-16']),
    ('UTF-32', ['UTF-32']),
    ('Q11 supplementary-plane counterexample', ['U+1F600']),
    ('Q13 mojibake bytes', ['E4 BD A0 E5 A5 BD', 'GBK']),
    ('Q14 input path', ['输入并显示']),
    ('Q14 storage/read path', ['保存并读取']),
]

QUESTION_HEADING_RE = re.compile(r'^#{1,3}\\s+(Q[0-9]+(?:-[A-Z])?)\\s+(.+?)\\s*$', re.M)
VISIBLE_Q_RE = re.compile(r'(?<![A-Za-z0-9])Q\\d+(?:-[A-Z])?(?![A-Za-z0-9])')

def inch(v):
    return float(v) / EMU_PER_INCH

def pt(v):
    return None if v is None else float(v) / 12700.0

def norm(s):
    s = s or ''
    s = s.replace('“','').replace('”','').replace('‘','').replace('’','')
    return re.sub(r'[^0-9A-Za-z\\u4e00-\\u9fff]+', '', s).lower()

def text_of_shape(shape):
    if not getattr(shape, 'has_text_frame', False):
        return ''
    try:
        return shape.text or ''
    except Exception:
        return ''

def iter_runs(shape):
    if not getattr(shape, 'has_text_frame', False):
        return
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.text and run.text.strip():
                yield run

def get_rgb(shape):
    try:
        fill = shape.fill
        if fill.type is None:
            return None
        rgb = fill.fore_color.rgb
        return str(rgb) if rgb else None
    except Exception:
        return None

def notes_for_slide(zf, slide_idx):
    rel_path = f'ppt/slides/_rels/slide{slide_idx}.xml.rels'
    if rel_path not in zf.namelist():
        return ''
    root = ET.fromstring(zf.read(rel_path))
    target = None
    for rel in root.findall('{http://schemas.openxmlformats.org/package/2006/relationships}Relationship'):
        if rel.attrib.get('Type','').endswith('/notesSlide'):
            target = rel.attrib.get('Target')
            break
    if not target:
        return ''
    base = posixpath.dirname(f'ppt/slides/slide{slide_idx}.xml')
    note_path = posixpath.normpath(posixpath.join(base, target))
    if note_path not in zf.namelist():
        return ''
    nroot = ET.fromstring(zf.read(note_path))
    texts = [e.text or '' for e in nroot.findall('.//a:t', NS)]
    return '\\n'.join(t.strip() for t in texts if t and t.strip())

def title_shape(slide):
    try:
        if slide.shapes.title is not None and text_of_shape(slide.shapes.title).strip():
            return slide.shapes.title
    except Exception:
        pass
    candidates=[]
    for sh in slide.shapes:
        txt=text_of_shape(sh).strip()
        if not txt:
            continue
        y=inch(sh.top)
        max_size=0
        for r in iter_runs(sh):
            sz=pt(r.font.size)
            if sz:
                max_size=max(max_size,sz)
        if 0.35 <= y <= 1.9:
            candidates.append((max_size, -y, -len(txt), sh))
    if not candidates:
        return None
    candidates.sort(reverse=True, key=lambda t:(t[0],t[1],t[2]))
    return candidates[0][3]

def shape_geom(sh):
    return [round(inch(sh.left),3), round(inch(sh.top),3), round(inch(sh.width),3), round(inch(sh.height),3)]

def dominant_font_and_size(sh):
    fonts=Counter(); sizes=Counter()
    for r in iter_runs(sh):
        fonts[r.font.name or '<inherited>'] += max(1,len(r.text.strip()))
        sz=pt(r.font.size)
        if sz is not None:
            sizes[round(sz,1)] += max(1,len(r.text.strip()))
    font=fonts.most_common(1)[0][0] if fonts else '<unknown>'
    size=sizes.most_common(1)[0][0] if sizes else None
    return font,size

def find_rails(slide, sw, sh):
    rails=[]
    for shape in slide.shapes:
        if get_rgb(shape) != EXPECTED['frame_violet']:
            continue
        x,y,w,h=map(inch,[shape.left,shape.top,shape.width,shape.height])
        if w >= sw-0.15 and h <= 0.18:
            rails.append((x,y,w,h))
    top=any(abs(x)<=0.05 and y<=0.14 and abs(h-EXPECTED['rail_h_in'])<=0.05 for x,y,w,h in rails)
    bottom=any(abs(x)<=0.05 and (y+h)>=sh-0.14 and abs(h-EXPECTED['rail_h_in'])<=0.05 for x,y,w,h in rails)
    return top,bottom,rails

def parse_questions(course_text):
    return [(qid, q.strip()) for qid,q in QUESTION_HEADING_RE.findall(course_text)]

def fuzzy_find_question(q, slide_titles):
    nq=norm(q)
    hits=[]
    if not nq:
        return hits
    for i,t in enumerate(slide_titles,1):
        nt=norm(t)
        if nq in nt or nt in nq:
            if min(len(nq),len(nt)) >= max(10, int(0.55*len(nq))):
                hits.append(i); continue
        if len(nq)>=12 and len(nt)>=12 and SequenceMatcher(None,nq,nt).ratio() >= 0.72:
            hits.append(i)
    return hits

def add_issue(issues, sev, category, msg, slide=None, detail=None):
    item={'severity':sev,'category':category,'message':msg}
    if slide is not None: item['slide']=slide
    if detail: item['detail']=detail
    issues.append(item)

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('deck')
    ap.add_argument('--course', required=True)
    ap.add_argument('--outdir', default='audit-output')
    args=ap.parse_args()
    deck=Path(args.deck); course=Path(args.course); outdir=Path(args.outdir); outdir.mkdir(parents=True,exist_ok=True)
    prs=Presentation(str(deck))
    sw,sh=inch(prs.slide_width), inch(prs.slide_height)
    course_text=course.read_text(encoding='utf-8')
    questions=parse_questions(course_text)
    issues=[]; slides=[]; all_visible=[]; font_counts=Counter(); size_counts=Counter(); title_geoms={}
    with zipfile.ZipFile(deck,'r') as zf:
        for idx,slide in enumerate(prs.slides,1):
            texts=[]
            for shape in slide.shapes:
                txt=text_of_shape(shape).strip()
                if txt: texts.append(txt)
                for r in iter_runs(shape):
                    font_counts[r.font.name or '<inherited>'] += max(1,len(r.text.strip()))
                    sz=pt(r.font.size)
                    if sz: size_counts[round(sz,1)] += max(1,len(r.text.strip()))
            visible='\\n'.join(texts)
            all_visible.append(visible)
            notes=notes_for_slide(zf,idx)
            ts=title_shape(slide)
            title=text_of_shape(ts).strip() if ts else ''
            tgeom=shape_geom(ts) if ts else None
            tfont,tsize=dominant_font_and_size(ts) if ts else ('<none>',None)
            title_geoms[idx]=tgeom
            toprail,bottomrail,rails=find_rails(slide,sw,sh)
            oob=[]
            for shobj in slide.shapes:
                x,y,w,h=map(inch,[shobj.left,shobj.top,shobj.width,shobj.height])
                if x < -0.02 or y < -0.02 or x+w > sw+0.02 or y+h > sh+0.02:
                    oob.append({'name':getattr(shobj,'name',''), 'geom':[round(x,3),round(y,3),round(w,3),round(h,3)]})
            small=[]
            for shobj in slide.shapes:
                txt=text_of_shape(shobj).strip()
                if not txt: continue
                run_sizes=[]
                for r in iter_runs(shobj):
                    sz=pt(r.font.size)
                    if sz: run_sizes.append(sz)
                if run_sizes and min(run_sizes) < 22 and len(txt) >= 35:
                    small.append({'text':txt[:90].replace('\\n',' / '), 'min_pt':round(min(run_sizes),1)})
            note_start='\\n'.join(notes.splitlines()[:10])
            notes_has_purpose=('[问题' in note_start or '[页面目的' in note_start)
            notes_has_transcript='[教师逐字稿]' in notes or '[逐字稿]' in notes
            slides.append({
                'slide':idx,'title':title,'title_geom':tgeom,'title_font':tfont,'title_size':tsize,
                'visible_chars':len(visible),'notes_chars':len(notes),'notes_has_purpose':notes_has_purpose,
                'notes_has_transcript':notes_has_transcript,'top_rail':toprail,'bottom_rail':bottomrail,
                'oob':oob,'small_text':small,'visible_q_labels':VISIBLE_Q_RE.findall(visible),
            })
            if idx>1 and (not toprail or not bottomrail):
                missing=('top' if not toprail else '')+(' and ' if (not toprail and not bottomrail) else '')+('bottom' if not bottomrail else '')
                add_issue(issues,'P1','frame',f'Missing required {missing} violet rail rectangle.',idx)
            if oob:
                add_issue(issues,'P1','geometry',f'{len(oob)} object(s) extend outside the slide canvas.',idx,detail=oob[:5])
            if idx>1 and ts:
                x,y,w,h=tgeom
                if abs(x-EXPECTED['title_x'])>0.08 or abs(y-EXPECTED['title_y'])>0.08 or abs(w-EXPECTED['title_w'])>0.20:
                    add_issue(issues,'P2','title geometry','Title does not match canonical baseline/width.',idx,detail={'geom':tgeom})
                if tsize is None or abs(tsize-EXPECTED['title_pt'])>1.0:
                    add_issue(issues,'P2','typography',f'Title size is {tsize} pt; expected 36 pt.',idx)
                if EXPECTED['title_font_fragment'].lower() not in (tfont or '').lower():
                    add_issue(issues,'P2','typography',f'Title dominant font is {tfont}; expected Alibaba PuHuiTi 3.0 115 Black.',idx)
            if idx>1 and not notes.strip():
                add_issue(issues,'P1','speaker notes','Speaker Notes are empty.',idx)
            elif idx>1:
                if not notes_has_purpose:
                    add_issue(issues,'P2','speaker notes','Notes do not begin with an explicit full question/page-purpose field.',idx)
                if len(notes)<160:
                    add_issue(issues,'P2','speaker notes',f'Notes are very short ({len(notes)} chars); transcript may not be classroom-usable.',idx)
                if not notes_has_transcript:
                    add_issue(issues,'P3','speaker notes','No explicit [教师逐字稿] marker; manually verify usable transcript.',idx)
            if VISIBLE_Q_RE.search(visible):
                add_issue(issues,'P1','student-facing content',f'Internal Q label(s) visible to students: {sorted(set(VISIBLE_Q_RE.findall(visible)))}',idx)
            if small:
                add_issue(issues,'P3','readability',f'{len(small)} longer text box(es) contain text below 22 pt; review whether legitimate micro/timeline/source text.',idx,detail=small[:4])

    if abs(sw-EXPECTED['width_in'])>0.03 or abs(sh-EXPECTED['height_in'])>0.03:
        add_issue(issues,'P0','canvas',f'Deck size is {sw:.3f} x {sh:.3f} in; expected 13.333 x 7.5 in.')

    slide_titles=[s['title'] for s in slides]
    qmap=[]
    for qid,q in questions:
        hits=fuzzy_find_question(q,slide_titles)
        qmap.append({'id':qid,'question':q,'slides':hits})
        if not hits:
            add_issue(issues,'P0','course alignment',f'{qid} question not found in slide titles: {q}')
        elif len(hits)==1:
            add_issue(issues,'P1','question-answer pairing',f'{qid} appears on only one slide; major questions should have question-only → duplicated answer/explanation slides.',hits[0])
        else:
            a,b=hits[0],hits[1]
            ga,gb=title_geoms.get(a),title_geoms.get(b)
            if ga and gb and any(abs(x-y)>0.02 for x,y in zip(ga,gb)):
                add_issue(issues,'P1','question-answer pairing',f'{qid} question geometry drifts between first two occurrences (slides {a} and {b}).',a,detail={'first':ga,'second':gb})

    alltext='\\n'.join(all_visible)
    for label,needles in TECH_REQUIRED:
        missing=[n for n in needles if n not in alltext]
        if missing:
            add_issue(issues,'P0','technical/course alignment',f'Missing required visible content for {label}: {missing}')
    if not any(k in alltext for k in ['编码单元','code unit','码元']):
        add_issue(issues,'P0','technical/course alignment','Q11-B does not visibly explain that 8/16/32 refer to encoding/code-unit width.')

    inherited_weight=font_counts.get('<inherited>',0)
    total_font_weight=sum(font_counts.values()) or 1
    if inherited_weight/total_font_weight > 0.05:
        add_issue(issues,'P2','typography',f'{inherited_weight/total_font_weight:.1%} of text-weight uses inherited/unspecified run font; style guide requires explicit Alibaba PuHuiTi 3.0.')
    bad_fonts=[(f,c) for f,c in font_counts.items() if f != '<inherited>' and 'Alibaba PuHuiTi 3.0' not in f]
    if bad_fonts:
        add_issue(issues,'P2','typography','Non-Alibaba native text fonts detected.',detail=sorted(bad_fonts,key=lambda x:-x[1])[:12])

    risk={1}
    for it in issues:
        if 'slide' in it and it['severity'] in ('P0','P1'):
            risk.add(it['slide'])
    for item in qmap:
        if item['id'] in ('Q10','Q11','Q11-A','Q11-B','Q14'):
            risk.update(item['slides'])
    risk=sorted(risk)[:12]

    counts=Counter(i['severity'] for i in issues)
    data={
        'deck':str(deck),'slides_count':len(slides),'slide_size_in':[round(sw,3),round(sh,3)],
        'font_counts':font_counts.most_common(), 'size_counts':size_counts.most_common(),
        'questions':qmap,'slides':slides,'issues':issues,'counts':dict(counts),'high_risk_slides':risk,
    }
    (outdir/'audit-data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')

    lines=['# PPTX strict audit report','',f'**Deck:** `{deck.name}`  ',f'**Slides:** {len(slides)}  ',f'**Canvas:** {sw:.3f} × {sh:.3f} in  ',f'**Course design:** `{course}`  ','',
           '## Summary','',f"- P0: {counts.get('P0',0)}",f"- P1: {counts.get('P1',0)}",f"- P2: {counts.get('P2',0)}",f"- P3: {counts.get('P3',0)}",'',
           '## Course-question coverage','', '| Question | Slides |','|---|---|']
    for q in qmap:
        lines.append(f"| {q['id']} — {q['question'].replace('|','/')} | {', '.join(map(str,q['slides'])) if q['slides'] else 'MISSING'} |")
    lines += ['','## Font summary','', '| Font | Weighted characters |','|---|---:|']
    for f,c in font_counts.most_common(12):
        lines.append(f"| {f.replace('|','/')} | {c} |")
    lines += ['','## Findings','']
    order={'P0':0,'P1':1,'P2':2,'P3':3}
    for i,it in enumerate(sorted(issues,key=lambda x:(order.get(x['severity'],9),x.get('slide',9999),x['category'])),1):
        where=f" — slide {it['slide']}" if 'slide' in it else ''
        lines.append(f"{i}. **{it['severity']} · {it['category']}**{where}: {it['message']}")
        if it.get('detail'):
            lines.append(f"   - Detail: `{json.dumps(it['detail'],ensure_ascii=False)[:1200]}`")
    if not issues:
        lines.append('No automated findings.')
    lines += ['','## Visual-review queue','',
              'High-risk slides selected for full-resolution review: ' + (', '.join(map(str,risk)) if risk else 'none'),'',
              '## Acceptance note','',
              'Automated checks do **not** certify WPS/PowerPoint rendering. A final WPS classroom check remains required by `slide-style-guide.md`. LibreOffice rendering in CI is used only to identify likely layout defects.','']
    (outdir/'audit-report.md').write_text('\\n'.join(lines),encoding='utf-8')
    print('\\n'.join(lines))

if __name__=='__main__':
    main()
